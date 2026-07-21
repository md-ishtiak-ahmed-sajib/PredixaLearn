from __future__ import annotations

import asyncio
import hashlib
import io
import json
import zipfile
from dataclasses import replace
from datetime import datetime, timedelta, timezone

import pytest
from docx import Document

from app.core.config import get_settings
from app.storage.history import HistoryStore
from app.teaching.batches import BatchCoordinator, build_batch_manifest
from app.teaching.service import (
    _bbox,
    _confidence,
    _shingles,
    corrected_export_path,
    duplicate_matches,
    effective_view,
    export_question_bank,
    parse_syllabus,
    parse_taxonomy,
    review_workspace,
    revision_pack_export_path,
    taxonomy_export,
)
from app.teaching.store import TeachingConflictError, TeachingStore, TeachingValidationError


def settings_for(tmp_path):
    return replace(
        get_settings(),
        output_dir=tmp_path / "output",
        upload_dir=tmp_path / "uploads",
        history_dir=tmp_path / "history",
        history_db_path=tmp_path / "history" / "history.sqlite3",
        log_dir=tmp_path / "logs",
    )


def completed(store: HistoryStore, job_id: str = "a" * 32) -> str:
    result = {
        "full_text": "1. Explain evidence. (5 marks)\n2. Compare evidence. (4 marks)",
        "lines": [
            {
                "line_id": "p1-l1",
                "page_number": 1,
                "text": "1. Explain evidence. (5 marks)",
                "confidence": 0.62,
                "normalized_bbox": [0.1, 0.1, 0.9, 0.18],
            },
            {
                "line_id": "p1-l2",
                "page_number": 1,
                "text": "2. Compare evidence. (4 marks)",
                "confidence": 0.94,
                "normalized_bbox": [0.1, 0.3, 0.9, 0.38],
            },
        ],
        "page_count": 1,
    }
    result_path = store.save_result(job_id, result)
    created = datetime(2026, 7, 21, tzinfo=timezone.utc)
    store.upsert(
        job_id=job_id,
        input_name="exam.pdf",
        workflow="text_recognition",
        input_kind="pdf",
        status="completed",
        item_count=1,
        error=None,
        result_path=result_path,
        language="en",
        document_profile="exam",
        created_at=created,
        started_at=created,
        completed_at=created + timedelta(seconds=2),
        updated_at=created + timedelta(seconds=2),
    )
    return job_id


def test_corrections_are_overlays_with_etags_and_hash_chained_audit(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    job_id = completed(store.history)
    correction = store.create_correction(
        job_id,
        {
            "target_kind": "line",
            "target_id": "p1-l1",
            "original": {"text": "1. Explain evidence. (5 marks)"},
            "replacement": {"text": "1. Explain the evidence. (5 marks)"},
            "reason": "Missing article visible in source",
            "status": "draft",
            "source_line_ids": ["p1-l1"],
            "geometry": [0.1, 0.1, 0.9, 0.18],
        },
    )
    with pytest.raises(TeachingConflictError):
        store.update_correction(correction["correction_id"], {"status": "approved"}, '"stale"')
    approved = store.update_correction(
        correction["correction_id"],
        {"status": "approved", "actor": "teacher-1"},
        correction["etag"],
    )
    result = store.history.read_result(store.history.get(job_id))
    effective = effective_view(result, store.list_corrections(job_id)["items"])

    assert result["lines"][0]["text"] == "1. Explain evidence. (5 marks)"
    assert effective["lines"][0]["text"] == "1. Explain the evidence. (5 marks)"
    assert approved["version"] == 2
    events = store.audit(job_id=job_id)
    assert events[1]["previous_hash"] == events[0]["event_hash"]


def test_review_workspace_exposes_heatmap_and_screen_reader_equivalent(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    job_id = completed(store.history)
    workspace = review_workspace(store, job_id)

    assert workspace["pages"][0]["lines"][0]["confidence_band"] == "low"
    assert {"text", "question", "marks", "unmapped"} <= set(
        workspace["pages"][0]["lines"][0]["evidence_types"]
    )
    assert workspace["pages"][0]["lines"][1]["confidence_band"] == "high"
    assert workspace["screen_reader_issues"][0]["line_id"] == "p1-l1"
    assert workspace["source_hash"] == store.source_hash(job_id)


def test_taxonomy_rejects_cycles_and_duplicate_aliases(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    with pytest.raises(TeachingValidationError, match="cycle"):
        store.create_taxonomy(
            {
                "name": "Cycle",
                "nodes": [
                    {"node_id": "a", "label": "A", "parent_id": "b"},
                    {"node_id": "b", "label": "B", "parent_id": "a"},
                ],
            }
        )
    with pytest.raises(TeachingValidationError, match="Duplicate"):
        store.create_taxonomy(
            {
                "name": "Aliases",
                "nodes": [
                    {"node_id": "a", "label": "Algebra", "aliases": ["Functions"]},
                    {"node_id": "b", "label": "Functions"},
                ],
            }
        )


def test_syllabus_import_and_publication_retains_source_checksum(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    data = b"objective_id,parent_id,code,label,description,aliases\nobj-1,,M1,Algebra,Use algebra,Equations|Functions\n"
    parsed = parse_syllabus(data, "syllabus.csv")
    saved = store.save_syllabus({**parsed, "name": "Math 2026", "status": "published"})

    assert saved["source_checksum"] == hashlib.sha256(data).hexdigest()
    assert saved["status"] == "published"
    assert saved["objectives"][0]["confirmed"] is True


def test_inferred_syllabus_requires_confirmation_and_publishes_new_version(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    parsed = {
        "name": "Inferred physics",
        "source_type": "docx",
        "source_checksum": "a" * 64,
        "objectives": [
            {
                "objective_id": "obj-1",
                "label": "Explain momentum",
                "inferred": True,
                "confirmed": False,
            }
        ],
    }
    draft = store.save_syllabus(parsed)
    with pytest.raises(TeachingValidationError, match="teacher-confirmed"):
        store.save_syllabus(
            {
                **parsed,
                "syllabus_id": draft["syllabus_id"],
                "status": "published",
            }
        )
    published = store.save_syllabus(
        {
            **parsed,
            "syllabus_id": draft["syllabus_id"],
            "status": "published",
            "objectives": [{**parsed["objectives"][0], "confirmed": True}],
        }
    )
    assert published["version"] == 2
    assert published["status"] == "published"
    assert store.list_syllabi()[0]["objectives"][0]["confirmed"] is True


def test_taxonomy_csv_json_round_trip_and_suggestion_mapping_gate(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    job_id = completed(store.history)
    parsed = parse_taxonomy(
        b"node_id,parent_id,label,aliases,description,sort_order\nroot,,Mechanics,Force|Motion,Physics,0\n",
        "topics.csv",
    )
    taxonomy = store.create_taxonomy({"name": "Physics", **parsed})
    json_body, media_type, _ = taxonomy_export(taxonomy, "json")
    assert media_type == "application/json"
    assert json.loads(json_body)["nodes"][0]["label"] == "Mechanics"

    syllabus = store.save_syllabus(
        {
            "name": "Physics syllabus",
            "source_type": "json",
            "source_checksum": "f" * 64,
            "objectives": [{"objective_id": "obj-1", "label": "Forces"}],
        }
    )
    with pytest.raises(TeachingValidationError, match="must remain drafts"):
        store.save_mapping(
            syllabus["syllabus_id"],
            {
                "job_id": job_id,
                "question_id": "q-1",
                "objective_id": "obj-1",
                "method": "ai_suggestion",
                "status": "approved",
            },
        )


def test_duplicate_detection_explains_numeric_variants_without_merging():
    matches = duplicate_matches(
        [
            {"question_id": "q1", "text": "Calculate force when mass is 10 kg and acceleration is 2."},
            {"question_id": "q2", "text": "Calculate force when mass is 12 kg and acceleration is 2."},
        ],
        threshold=0.4,
    )
    assert matches[0]["changed_numbers"] is True
    assert matches[0]["automatic_action"] == "none"
    assert matches[0]["shared_phrases"]


def test_question_bank_export_is_approved_checksums_and_no_source_images(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    job_id = completed(store.history)
    item = store.upsert_bank_item(
        {
            "job_id": job_id,
            "question_id": "q-1",
            "status": "approved",
            "actor": "teacher-1",
            "payload": {
                "text": "Explain the evidence.",
                "marks": 5,
                "source_job_id": job_id,
                "source_page": 1,
                "language": "en",
            },
        }
    )
    exported = export_question_bank(store, formats=["qti", "csv", "jsonl"])
    archive = store.settings.history_dir / "question-bank-exports" / exported["export_id"] / "predixalearn-question-bank.zip"

    assert item["status"] == "approved"
    assert exported["source_images_requested"] is False
    assert exported["source_images_included"] == []
    assert archive.is_file()
    assert json.loads((archive.parent / "manifest.json").read_text())["item_count"] == 1


def test_revision_pack_exports_student_safe_markdown_json_and_docx(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    job_id = completed(store.history)
    item = store.upsert_bank_item(
        {
            "job_id": job_id,
            "question_id": "q-1",
            "status": "approved",
            "payload": {
                "text": "Explain evidence.",
                "topic": "Methods",
                "difficulty": "medium",
                "teacher_notes": "Never export this",
            },
        }
    )
    pack = store.save_revision_pack(
        {
            "title": "Evidence revision",
            "status": "approved",
            "source_item_ids": [item["item_id"]],
            "payload": {"introduction": "Practice verified questions."},
        }
    )
    for format_name in ("markdown", "json", "docx"):
        assert revision_pack_export_path(store, pack["pack_id"], format_name).is_file()
    exported = json.loads(
        revision_pack_export_path(store, pack["pack_id"], "json").read_text(
            encoding="utf-8"
        )
    )
    assert exported["student_safe"] is True
    assert "teacher_notes" not in json.dumps(exported)


def test_batch_records_recover_interrupted_items_without_losing_staged_source(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    root = store.settings.upload_dir / "batches"
    root.mkdir(parents=True)
    staged = root / "sample.png"
    staged.write_bytes(b"source")
    batch = store.create_batch(
        {"name": "Offline", "workflow": "text_recognition", "language": "en"},
        [
            {
                "input_name": "sample.png",
                "staged_path": "sample.png",
                "source_sha256": hashlib.sha256(b"source").hexdigest(),
                "media_type": "image/png",
                "size_bytes": 6,
            }
        ],
    )
    store.patch_batch_item(batch["batch_id"], batch["items"][0]["item_id"], {"status": "processing"})
    assert store.recover_batches() == 1
    recovered = store.get_batch(batch["batch_id"])["items"][0]
    assert recovered["status"] == "queued"
    assert recovered["retry_pending"] == 1
    assert staged.is_file()


def test_teaching_store_rejects_new_non_english_content(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    with pytest.raises(TeachingValidationError, match="English revision packs only"):
        store.save_revision_pack({"title": "French", "language": "fr"})

    staged = store.settings.upload_dir / "batches" / "sample.png"
    staged.parent.mkdir(parents=True)
    staged.write_bytes(b"source")
    entry = {
        "input_name": "sample.png",
        "staged_path": "sample.png",
        "source_sha256": hashlib.sha256(b"source").hexdigest(),
        "media_type": "image/png",
        "size_bytes": 6,
    }
    with pytest.raises(TeachingValidationError, match="English batch documents only"):
        store.create_batch(
            {"name": "French", "workflow": "text_recognition", "language": "fr"},
            [entry],
        )
    with pytest.raises(TeachingValidationError, match="overrides must use English"):
        store.create_batch(
            {"name": "Override", "workflow": "text_recognition", "language": "en"},
            [{**entry, "overrides": {"language": "ar"}}],
        )


def test_batch_per_file_overrides_reorder_and_manifest_exclude_sources(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    root = store.settings.upload_dir / "batches"
    root.mkdir(parents=True)
    first, second = root / "one.png", root / "two.png"
    first.write_bytes(b"first-source")
    second.write_bytes(b"second-source")
    batch = store.create_batch(
        {"name": "English documents", "workflow": "text_recognition", "language": "en"},
        [
            {"input_name": "one.png", "staged_path": "one.png", "source_sha256": hashlib.sha256(first.read_bytes()).hexdigest(), "media_type": "image/png", "size_bytes": first.stat().st_size},
            {"input_name": "two.png", "staged_path": "two.png", "source_sha256": hashlib.sha256(second.read_bytes()).hexdigest(), "media_type": "image/png", "size_bytes": second.stat().st_size, "overrides": {"document_profile": "exam"}},
        ],
    )
    moved = store.reorder_batch_item(batch["batch_id"], batch["items"][1]["item_id"], direction=-1)
    assert moved["position"] == 0
    assert moved["overrides"] == {"document_profile": "exam"}
    archive = build_batch_manifest(store, batch["batch_id"])
    assert archive.is_file()
    with zipfile.ZipFile(archive) as bundle:
        assert set(bundle.namelist()) == {"manifest.json", "SHA256SUMS.txt"}
        manifest = json.loads(bundle.read("manifest.json"))
    assert manifest["staged_sources_included"] is False
    assert first.is_file() and second.is_file()

    prioritized = store.prioritize_batch_item(
        batch["batch_id"], batch["items"][0]["item_id"]
    )
    assert prioritized["position"] == 0


def test_syllabus_docx_case_json_and_import_validation_paths(tmp_path):
    document = Document()
    document.add_paragraph("Explain conservation of momentum")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Apply Newton's laws"
    table.cell(0, 1).text = "Evaluate force diagrams"
    stream = io.BytesIO()
    document.save(stream)

    parsed_docx = parse_syllabus(stream.getvalue(), "syllabus.docx")
    assert parsed_docx["source_type"] == "docx"
    assert parsed_docx["objectives"][0]["inferred"] is True
    assert parsed_docx["warnings"]

    case = {
        "CFItems": [
            {
                "identifier": "case-1",
                "humanCodingScheme": "PHY-1",
                "fullStatement": "Explain momentum",
            }
        ]
    }
    parsed_case = parse_syllabus(json.dumps(case).encode(), "framework.case")
    assert parsed_case["source_type"] == "ims_case_json"
    assert parsed_case["objectives"][0]["objective_id"] == "case-1"

    parsed_json = parse_syllabus(
        json.dumps({"objectives": [{"objective_id": "j-1", "label": "Energy"}]}).encode(),
        "framework.json",
    )
    assert parsed_json["source_type"] == "json"
    with pytest.raises(ValueError, match="Supported syllabus"):
        parse_syllabus(b"text", "framework.txt")
    with pytest.raises(ValueError, match="No syllabus"):
        parse_syllabus(b"objective_id,label\n", "empty.csv")
    with pytest.raises(ValueError, match="nodes array"):
        parse_taxonomy(b'{"nodes":"invalid"}', "invalid.json")
    with pytest.raises(ValueError, match="supports CSV"):
        parse_taxonomy(b"x", "invalid.txt")
    with pytest.raises(ValueError, match="contains no"):
        parse_taxonomy(b"node_id,label\n", "empty.csv")
    with pytest.raises(ValueError, match="format is invalid"):
        taxonomy_export({}, "xml")


def test_effective_question_and_document_overlays_and_geometry_guards(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    job_id = completed(store.history)
    for kind, target, replacement in (
        ("question", "q-1", {"text": "Explain the verified evidence."}),
        ("marks", "q-1", {"marks": 7}),
        ("reconstructed_text", "document", {"text": "Teacher-approved full document"}),
    ):
        correction = store.create_correction(
            job_id,
            {
                "target_kind": kind,
                "target_id": target,
                "replacement": replacement,
                "reason": "Source verified",
            },
        )
        store.update_correction(correction["correction_id"], {"status": "approved"}, correction["etag"])

    workspace = review_workspace(store, job_id)
    assert workspace["reconstructed_text"] == "Teacher-approved full document"
    assert workspace["questions"][0]["marks"] == 7
    assert workspace["questions"][0]["text"] == "Explain the verified evidence."
    assert _confidence({"retry_confidence": "bad", "score": 0.8}) == 0.8
    assert _confidence({"confidence": 2}) is None
    assert _bbox({"normalized_bbox": "bad"}) is None
    assert _bbox({"normalized_bbox": [0, "x", 1, 1]}) is None
    assert _bbox({"normalized_bbox": [0.9, 0, 0.1, 1]}) is None
    assert _shingles("") == set()
    assert _shingles("one two") == {("one", "two")}

    with pytest.raises(TeachingValidationError, match="unknown OCR"):
        store.create_correction(job_id, {"target_kind": "line", "target_id": "missing", "replacement": "x", "reason": "test"})
    with pytest.raises(TeachingValidationError, match="normalized"):
        store.create_correction(job_id, {"target_kind": "table_cell", "target_id": "c1", "source_line_ids": ["p1-l1"], "geometry": [0.8, 0, 0.2, 1], "replacement": "x", "reason": "test"})
    with pytest.raises(TeachingValidationError, match="known source"):
        store.create_correction(job_id, {"target_kind": "table_cell", "target_id": "c1", "source_line_ids": ["missing"], "replacement": "x", "reason": "test"})
    with pytest.raises(TeachingValidationError, match="document correction"):
        store.create_correction(job_id, {"target_kind": "reconstructed_text", "target_id": "page", "replacement": "x", "reason": "test"})
    with pytest.raises(ValueError, match="format is invalid"):
        corrected_export_path(store, job_id, "pdf")


def test_legacy_topics_migrate_non_destructively(tmp_path):
    settings = settings_for(tmp_path)
    history = HistoryStore(settings)
    job_id = completed(history)
    sidecar = settings.history_dir / "jobs" / job_id / "education" / "analysis.json"
    sidecar.parent.mkdir(parents=True, exist_ok=True)
    sidecar.write_text(
        json.dumps(
            {
                "topic_taxonomy": ["Mechanics"],
                "questions": [
                    {"teacher_override": {"topic": "Energy"}, "ai": {"topic": "Motion"}}
                ],
            }
        ),
        encoding="utf-8",
    )
    store = TeachingStore(settings)
    taxonomy = store.get_taxonomy("legacy-topics")
    labels = {node["label"] for node in taxonomy["nodes"]}
    assert {"Unmapped legacy topics", "Mechanics", "Energy", "Motion"} <= labels
    assert json.loads(sidecar.read_text(encoding="utf-8"))["topic_taxonomy"] == ["Mechanics"]


def test_batch_coordinator_safety_and_failure_paths(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    root = store.settings.upload_dir / "batches"
    root.mkdir(parents=True)
    source = root / "broken.png"
    source.write_bytes(b"actual")
    batch = store.create_batch(
        {"name": "Failure", "workflow": "text_recognition", "language": "en"},
        [
            {
                "input_name": "broken.png",
                "staged_path": "broken.png",
                "source_sha256": hashlib.sha256(b"different").hexdigest(),
                "media_type": "image/png",
                "size_bytes": source.stat().st_size,
            }
        ],
    )
    coordinator = BatchCoordinator(store)
    candidate = coordinator._next_item()
    assert candidate is not None
    asyncio.run(coordinator._process(*candidate))
    failed = store.get_batch(batch["batch_id"])["items"][0]
    assert failed["status"] == "failed"
    assert failed["retry_pending"] == 1
    with pytest.raises(ValueError, match="outside"):
        coordinator._safe_staged_path(str(tmp_path / "escape.png"))

    store.patch_batch_item(batch["batch_id"], failed["item_id"], {"status": "cancelled"})
    store.patch_batch(batch["batch_id"], {"status": "processing"})
    assert coordinator._next_item() is None
    assert store.get_batch(batch["batch_id"])["status"] == "partial"
    async def exercise_lifecycle() -> None:
        coordinator.start()
        coordinator.start()
        await coordinator.stop()

    asyncio.run(exercise_lifecycle())


def test_batch_coordinator_stops_a_legacy_non_english_override(tmp_path):
    store = TeachingStore(settings_for(tmp_path))
    root = store.settings.upload_dir / "batches"
    root.mkdir(parents=True)
    source = root / "legacy.png"
    source.write_bytes(b"legacy-source")
    batch = store.create_batch(
        {"name": "Legacy", "workflow": "text_recognition", "language": "en"},
        [
            {
                "input_name": "legacy.png",
                "staged_path": "legacy.png",
                "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                "media_type": "image/png",
                "size_bytes": source.stat().st_size,
            }
        ],
    )
    # Simulate an item created by an earlier multilingual-capable build. New
    # writes cannot create this state, but restart recovery must still be safe.
    with store.connect() as connection:
        connection.execute(
            "UPDATE batch_items SET overrides_json = ? WHERE item_id = ?",
            (json.dumps({"language": "ar"}), batch["items"][0]["item_id"]),
        )

    coordinator = BatchCoordinator(store)
    candidate = coordinator._next_item()
    assert candidate is not None
    asyncio.run(coordinator._process(*candidate))
    failed = store.get_batch(batch["batch_id"])["items"][0]
    assert failed["status"] == "failed"
    assert failed["retry_pending"] == 1
    assert "English-only release" in failed["error"]
    assert source.is_file()

"""Past Paper Intelligence persistence and grounded OpenAI enrichment."""

from __future__ import annotations

import hashlib
import json
import os
import re
from collections import Counter, defaultdict
from collections.abc import Mapping, Sequence
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Protocol

from docx import Document
from PIL import Image

from app.core.serialization import atomic_write_json, atomic_write_text
from app.education.questions import segment_questions
from app.storage.history import HistoryRecord, HistoryStore

_QUESTION_ID = re.compile(r"^q-[a-z0-9-]{1,80}$")
_REVIEW_STATUSES = {"needs_teacher_review", "teacher_reviewed", "approved"}
_DIFFICULTIES = {"low", "medium", "high", "unclassified"}
_MAX_TAXONOMY_ITEMS = 100
_MAX_TAXONOMY_ITEM_LENGTH = 120


class EducationError(RuntimeError):
    """A safe user-facing Past Paper Intelligence error."""


class OpenAIResponsesClient(Protocol):
    def create(self, **kwargs: Any) -> Any: ...


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _analysis_root(store: HistoryStore, record: HistoryRecord) -> Path:
    root = store.settings.history_dir / "jobs" / record.job_id / "education"
    root.mkdir(parents=True, exist_ok=True)
    return root


def _analysis_path(store: HistoryStore, record: HistoryRecord) -> Path:
    return _analysis_root(store, record) / "analysis.json"


def _result_hash(result: Any) -> str:
    payload = json.dumps(result, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _normalise_taxonomy(value: Sequence[str] | None) -> list[str]:
    if not value:
        return []
    normalized: list[str] = []
    for item in value:
        text = " ".join(str(item).split()).strip()
        if not text:
            continue
        if len(text) > _MAX_TAXONOMY_ITEM_LENGTH:
            raise EducationError("A topic taxonomy item is too long")
        if text.casefold() not in {candidate.casefold() for candidate in normalized}:
            normalized.append(text)
    if len(normalized) > _MAX_TAXONOMY_ITEMS:
        raise EducationError("A topic taxonomy can contain at most 100 entries")
    return normalized


def _base_analysis(
    result: Mapping[str, Any],
    *,
    job_id: str,
    subject: str | None,
    academic_level: str | None,
    taxonomy: Sequence[str] | None,
) -> dict[str, Any]:
    segmented = segment_questions(result)
    questions = [dict(item) for item in segmented["questions"]]
    for question in questions:
        question["ai"] = {
            "status": "not_requested",
            "topic": None,
            "subtopic": None,
            "question_type": None,
            "difficulty": "unclassified",
            "difficulty_reason": None,
            "cognitive_skill": None,
            "classification_confidence": None,
            "marks_inference": None,
        }
        question["teacher_override"] = {}
    return {
        "version": 1,
        "document_id": job_id,
        "ocr_result_sha256": _result_hash(result),
        "created_at": _now(),
        "updated_at": _now(),
        "subject": " ".join((subject or "").split()) or None,
        "academic_level": " ".join((academic_level or "").split()) or None,
        "topic_taxonomy": _normalise_taxonomy(taxonomy),
        "questions": questions,
        "segmentation_warnings": segmented["segmentation_warnings"],
        "analysis": _aggregate_questions(questions),
        "model": {
            "status": "not_requested",
            "provider": "openai",
            "model": None,
            "fallback_used": False,
            "consent_to_cloud": False,
        },
        "limitations": [
            "Question boundaries and marks are derived from OCR evidence and may need review.",
            "Historical patterns help prioritise revision; they do not predict exact future exam questions.",
        ],
    }


def _effective_value(question: Mapping[str, Any], field: str) -> Any:
    teacher = question.get("teacher_override")
    if isinstance(teacher, Mapping) and field in teacher:
        return teacher[field]
    ai = question.get("ai")
    return ai.get(field) if isinstance(ai, Mapping) else None


def _aggregate_questions(questions: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    topic_counts: Counter[str] = Counter()
    marks_by_topic: Counter[str] = Counter()
    difficulty_counts: Counter[str] = Counter()
    total_marks = 0
    for question in questions:
        topic = _effective_value(question, "topic")
        normalized_topic = str(topic).strip() if topic else "Unclassified"
        topic_counts[normalized_topic] += 1
        # A teacher's explicit marks correction is authoritative for all local
        # summaries.  The OCR value remains untouched alongside it as evidence.
        teacher = question.get("teacher_override")
        marks = (
            teacher.get("marks")
            if isinstance(teacher, Mapping) and "marks" in teacher
            else question.get("marks")
        )
        if isinstance(marks, int):
            marks_by_topic[normalized_topic] += marks
            total_marks += marks
        difficulty = _effective_value(question, "difficulty") or "unclassified"
        difficulty_counts[str(difficulty)] += 1
    distributions = [
        {"topic": topic, "question_count": count, "marks": marks_by_topic[topic]}
        for topic, count in topic_counts.most_common()
    ]
    priorities = [
        {
            "topic": item["topic"],
            "reason": f"Appears in {item['question_count']} question(s) worth {item['marks']} explicit mark(s).",
            "source_question_ids": [
                str(question.get("question_id"))
                for question in questions
                if (_effective_value(question, "topic") or "Unclassified") == item["topic"]
            ],
        }
        for item in distributions
        if item["topic"] != "Unclassified"
    ]
    return {
        "question_count": len(questions),
        "explicit_marks_total": total_marks,
        "topic_distribution": distributions,
        "marks_distribution": [
            {"topic": item["topic"], "marks": item["marks"]} for item in distributions
        ],
        "difficulty_distribution": [
            {"difficulty": name, "question_count": count}
            for name, count in sorted(difficulty_counts.items())
        ],
        "recurring_topics": [],
        "revision_priorities": priorities,
    }


def _model_schema() -> dict[str, Any]:
    question = {
        "type": "object",
        "additionalProperties": False,
        "required": [
            "question_id",
            "source_page",
            "source_line_ids",
            "topic",
            "subtopic",
            "question_type",
            "difficulty",
            "difficulty_reason",
            "cognitive_skill",
            "classification_confidence",
            "marks_inference",
        ],
        "properties": {
            "question_id": {"type": "string"},
            "source_page": {"type": "integer"},
            "source_line_ids": {"type": "array", "items": {"type": "string"}},
            "topic": {"type": ["string", "null"]},
            "subtopic": {"type": ["string", "null"]},
            "question_type": {"type": ["string", "null"]},
            "difficulty": {"type": "string", "enum": ["low", "medium", "high", "unclassified"]},
            "difficulty_reason": {"type": ["string", "null"]},
            "cognitive_skill": {"type": ["string", "null"]},
            "classification_confidence": {"type": ["number", "null"], "minimum": 0, "maximum": 1},
            "marks_inference": {
                "type": ["object", "null"],
                "additionalProperties": False,
                "required": ["marks", "reason"],
                "properties": {
                    "marks": {"type": ["integer", "null"], "minimum": 0, "maximum": 100},
                    "reason": {"type": ["string", "null"]},
                },
            },
        },
    }
    return {
        "type": "object",
        "additionalProperties": False,
        "required": ["questions", "practice_items"],
        "properties": {
            "questions": {"type": "array", "items": question},
            "practice_items": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "required": ["source_question_id", "prompt", "review_status"],
                    "properties": {
                        "source_question_id": {"type": "string"},
                        "prompt": {"type": "string"},
                        "review_status": {"type": "string", "enum": ["needs_teacher_review"]},
                    },
                },
            },
        },
    }


def _model_input(analysis: Mapping[str, Any]) -> list[dict[str, Any]]:
    questions = []
    for item in analysis.get("questions", []):
        if not isinstance(item, Mapping):
            continue
        questions.append(
            {
                "question_id": item.get("question_id"),
                "text": item.get("text"),
                "source_page": item.get("source_page"),
                "source_line_ids": item.get("source_line_ids"),
                "marks": item.get("marks"),
                "ocr_confidence": item.get("ocr_confidence"),
            }
        )
    instructions = (
        "You are PredixaLearn Past Paper Intelligence. Classify only the supplied OCR questions. "
        "Do not repair missing text, invent questions, predict future exams, or cite evidence that is not "
        "provided. Preserve exact question_id, source_page, and source_line_ids. Treat topics, difficulty, "
        "and any inferred marks as reviewable suggestions. Practice prompts must be clearly teacher-reviewable."
    )
    payload = {
        "subject": analysis.get("subject"),
        "academic_level": analysis.get("academic_level"),
        "topic_taxonomy": analysis.get("topic_taxonomy"),
        "questions": questions,
    }
    return [
        {"role": "developer", "content": instructions},
        {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
    ]


def _response_json(response: Any) -> dict[str, Any]:
    text = getattr(response, "output_text", None)
    if not isinstance(text, str) or not text.strip():
        raise EducationError("The analysis service returned no structured output")
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise EducationError("The analysis service returned invalid structured output") from exc
    if not isinstance(parsed, dict):
        raise EducationError("The analysis service returned an invalid analysis shape")
    return parsed


def _validate_model_output(analysis: dict[str, Any], raw: Mapping[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    expected = {
        str(item["question_id"]): item
        for item in analysis["questions"]
        if isinstance(item, Mapping) and item.get("question_id")
    }
    items = raw.get("questions")
    if not isinstance(items, Sequence) or isinstance(items, (str, bytes)):
        raise EducationError("The analysis service omitted question classifications")
    validated: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in items:
        if not isinstance(item, Mapping):
            raise EducationError("The analysis service returned a malformed question classification")
        identifier = str(item.get("question_id", ""))
        source = expected.get(identifier)
        if source is None or identifier in seen:
            raise EducationError("The analysis service referenced an unknown question")
        seen.add(identifier)
        if int(item.get("source_page", -1)) != int(source.get("source_page", -2)):
            raise EducationError("The analysis service returned an invalid source page")
        source_lines = {str(value) for value in source.get("source_line_ids", [])}
        returned_lines = [str(value) for value in item.get("source_line_ids", [])]
        if not returned_lines or not set(returned_lines).issubset(source_lines):
            raise EducationError("The analysis service returned invalid source line references")
        confidence = item.get("classification_confidence")
        if confidence is not None and not isinstance(confidence, (int, float)):
            raise EducationError("The analysis service returned an invalid AI confidence")
        difficulty = str(item.get("difficulty", "unclassified"))
        if difficulty not in _DIFFICULTIES:
            raise EducationError("The analysis service returned an invalid difficulty")
        validated.append(
            {
                "question_id": identifier,
                "status": "suggested",
                "topic": item.get("topic"),
                "subtopic": item.get("subtopic"),
                "question_type": item.get("question_type"),
                "difficulty": difficulty,
                "difficulty_reason": item.get("difficulty_reason"),
                "cognitive_skill": item.get("cognitive_skill"),
                "classification_confidence": round(float(confidence), 4) if confidence is not None else None,
                "marks_inference": item.get("marks_inference"),
                "source_page": int(item["source_page"]),
                "source_line_ids": returned_lines,
            }
        )
    practice = raw.get("practice_items", [])
    if not isinstance(practice, Sequence) or isinstance(practice, (str, bytes)):
        raise EducationError("The analysis service returned malformed practice suggestions")
    validated_practice: list[dict[str, Any]] = []
    for item in practice[:20]:
        if not isinstance(item, Mapping):
            raise EducationError("The analysis service returned a malformed practice suggestion")
        source_id = str(item.get("source_question_id", ""))
        prompt = " ".join(str(item.get("prompt", "")).split())
        if source_id not in expected or not prompt or len(prompt) > 1200:
            raise EducationError("The analysis service returned an ungrounded practice suggestion")
        validated_practice.append(
            {
                "source_question_id": source_id,
                "prompt": prompt,
                "review_status": "needs_teacher_review",
            }
        )
    return validated, validated_practice


def _get_openai_client() -> OpenAIResponsesClient:
    if not os.getenv("OPENAI_API_KEY", "").strip():
        raise EducationError("OpenAI analysis is unavailable. Set OPENAI_API_KEY and restart PredixaLearn.")
    try:
        from openai import OpenAI
    except ImportError as exc:  # pragma: no cover - caught by package installation tests
        raise EducationError("OpenAI analysis support is not installed in this runtime") from exc
    return OpenAI().responses


def _call_model(client: OpenAIResponsesClient, model: str, analysis: Mapping[str, Any]) -> dict[str, Any]:
    response = client.create(
        model=model,
        input=_model_input(analysis),
        reasoning={"effort": "medium"},
        text={
            "format": {
                "type": "json_schema",
                "name": "past_paper_intelligence",
                "strict": True,
                "schema": _model_schema(),
            }
        },
    )
    return _response_json(response)


def _enrich(analysis: dict[str, Any], client: OpenAIResponsesClient | None = None) -> None:
    responses = client or _get_openai_client()
    fallback_used = False
    try:
        raw = _call_model(responses, "gpt-5.6-sol", analysis)
        model = "gpt-5.6-sol"
    except Exception as primary_error:
        try:
            raw = _call_model(responses, "gpt-5.6-terra", analysis)
            model = "gpt-5.6-terra"
            fallback_used = True
        except Exception as fallback_error:
            raise EducationError("OpenAI analysis is temporarily unavailable; no OCR data was changed.") from fallback_error
        del primary_error
    suggestions, practice = _validate_model_output(analysis, raw)
    by_id = {str(item["question_id"]): item for item in analysis["questions"]}
    for suggestion in suggestions:
        by_id[str(suggestion["question_id"])]["ai"] = suggestion
    analysis["practice_items"] = practice
    analysis["model"] = {
        "status": "completed",
        "provider": "openai",
        "model": model,
        "fallback_used": fallback_used,
        "consent_to_cloud": True,
    }
    analysis["analysis"] = _aggregate_questions(analysis["questions"])


def _write_docx(path: Path, analysis: Mapping[str, Any]) -> None:
    document = Document()
    document.add_heading("PredixaLearn Past Paper Intelligence", level=0)
    document.add_paragraph("Teacher-reviewable, source-linked analysis. It does not predict future exams.")
    document.add_paragraph(f"Document ID: {analysis.get('document_id')}")
    for question in analysis.get("questions", []):
        if not isinstance(question, Mapping):
            continue
        document.add_heading(str(question.get("display_number", "Question")), level=1)
        document.add_paragraph(str(question.get("text", "")))
        ai = question.get("ai") if isinstance(question.get("ai"), Mapping) else {}
        document.add_paragraph(
            f"Source: page {question.get('source_page')} · lines {', '.join(question.get('source_line_ids', []))}"
        )
    practice = analysis.get("practice_items")
    if isinstance(practice, Sequence) and not isinstance(practice, (str, bytes)) and practice:
        document.add_heading("Teacher-reviewable practice suggestions", level=1)
        for item in practice:
            if not isinstance(item, Mapping):
                continue
            document.add_paragraph(
                f"From {item.get('source_question_id')}: {item.get('prompt', '')} "
                "(teacher review required)"
            )
        document.add_paragraph(
            f"Topic: {_effective_value(question, 'topic') or 'Unclassified'} · "
            f"Difficulty: {_effective_value(question, 'difficulty') or 'Unclassified'} · "
            f"OCR confidence: {question.get('ocr_confidence')} · "
            f"AI confidence: {ai.get('classification_confidence')}"
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _analysis_markdown(analysis: Mapping[str, Any]) -> str:
    lines = [
        "# PredixaLearn Past Paper Intelligence",
        "",
        "> Teacher-reviewable, source-linked analysis. It analyzes historical patterns and does not predict exact future questions.",
        "",
        f"- Document: `{analysis.get('document_id')}`",
        f"- Subject: {analysis.get('subject') or 'Not specified'}",
        f"- Academic level: {analysis.get('academic_level') or 'Not specified'}",
        "",
        "## Questions",
        "",
    ]
    for question in analysis.get("questions", []):
        if not isinstance(question, Mapping):
            continue
        ai = question.get("ai") if isinstance(question.get("ai"), Mapping) else {}
        lines.extend(
            [
                f"### {question.get('display_number')}",
                "",
                str(question.get("text", "")),
                "",
                f"Source: page {question.get('source_page')} · lines {', '.join(question.get('source_line_ids', []))}",
                f"OCR confidence: {question.get('ocr_confidence') if question.get('ocr_confidence') is not None else 'not available'}",
                f"Topic: {_effective_value(question, 'topic') or 'Unclassified'}",
                f"Difficulty: {_effective_value(question, 'difficulty') or 'Unclassified'}",
                f"AI confidence: {ai.get('classification_confidence') if ai.get('classification_confidence') is not None else 'not available'}",
                "",
            ]
        )
    practice = analysis.get("practice_items")
    if isinstance(practice, Sequence) and not isinstance(practice, (str, bytes)) and practice:
        lines.extend(["## Teacher-reviewable practice suggestions", ""])
        lines.extend(
            f"- From `{item.get('source_question_id')}`: {item.get('prompt', '')} "
            "*(teacher review required)*"
            for item in practice
            if isinstance(item, Mapping)
        )
        lines.append("")
    return "\n".join(lines)


def _write_analysis_files(store: HistoryStore, record: HistoryRecord, analysis: dict[str, Any]) -> None:
    root = _analysis_root(store, record)
    analysis["updated_at"] = _now()
    atomic_write_json(root / "analysis.json", analysis)
    atomic_write_text(root / "analysis.md", _analysis_markdown(analysis))
    _write_docx(root / "analysis.docx", analysis)


def _append_audit(store: HistoryStore, record: HistoryRecord, event: Mapping[str, Any]) -> None:
    audit = _analysis_root(store, record) / "audit.jsonl"
    entry = {"at": _now(), **dict(event)}
    with audit.open("a", encoding="utf-8", newline="\n") as handle:
        handle.write(json.dumps(entry, ensure_ascii=False, sort_keys=True) + "\n")


def _build_question_crops(store: HistoryStore, record: HistoryRecord, analysis: Mapping[str, Any]) -> None:
    manifest = store.read_manifest(record)
    if not manifest:
        return
    archive_relative = manifest.get("archive_root")
    if not isinstance(archive_relative, str):
        return
    archive_root = store._resolve_history_path(archive_relative)  # noqa: SLF001
    pages_root = archive_root / "source-pages"
    crops_root = _analysis_root(store, record) / "crops"
    crops_root.mkdir(parents=True, exist_ok=True)
    for question in analysis.get("questions", []):
        if not isinstance(question, Mapping):
            continue
        identifier = str(question.get("question_id", ""))
        bbox = question.get("source_bbox")
        page = question.get("source_page")
        if not _QUESTION_ID.fullmatch(identifier) or not isinstance(page, int):
            continue
        if not isinstance(bbox, Sequence) or isinstance(bbox, (str, bytes)) or len(bbox) != 4:
            continue
        page_path = pages_root / f"page-{page:03d}.png"
        if not page_path.is_file():
            continue
        try:
            with Image.open(page_path) as source:
                image = source.convert("RGB")
                left, top, right, bottom = [float(value) for value in bbox]
                width, height = image.size
                crop = image.crop(
                    (
                        max(0, int(left * width) - 12),
                        max(0, int(top * height) - 12),
                        min(width, int(right * width) + 12),
                        min(height, int(bottom * height) + 12),
                    )
                )
                if crop.width > 1 and crop.height > 1:
                    crop.save(crops_root / f"{identifier}.png", format="PNG", optimize=True)
        except (OSError, TypeError, ValueError):
            continue


def read_analysis(store: HistoryStore, record: HistoryRecord, *, create_preview: bool = True) -> dict[str, Any]:
    result = store.read_result(record)
    if not isinstance(result, Mapping):
        raise EducationError("This OCR result cannot be analyzed")
    path = _analysis_path(store, record)
    current_hash = _result_hash(result)
    if path.is_file():
        try:
            analysis = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, UnicodeError, json.JSONDecodeError) as exc:
            raise EducationError("The saved Past Paper Intelligence record is corrupted") from exc
        if isinstance(analysis, dict) and analysis.get("ocr_result_sha256") == current_hash:
            return analysis
    if create_preview:
        return _base_analysis(result, job_id=record.job_id, subject=None, academic_level=None, taxonomy=None)
    raise EducationError("No Past Paper Intelligence analysis has been saved for this run")


def create_analysis(
    store: HistoryStore,
    record: HistoryRecord,
    *,
    subject: str | None,
    academic_level: str | None,
    taxonomy: Sequence[str] | None,
    consent_to_cloud: bool,
    client: OpenAIResponsesClient | None = None,
) -> dict[str, Any]:
    result = store.read_result(record)
    if not isinstance(result, Mapping):
        raise EducationError("This OCR result cannot be analyzed")
    previous: dict[str, Any] | None = None
    try:
        previous = read_analysis(store, record, create_preview=False)
    except EducationError:
        # A first analysis legitimately has no sidecar. A corrupt or stale
        # sidecar is never trusted over the current immutable OCR result.
        previous = None
    analysis = _base_analysis(
        result,
        job_id=record.job_id,
        subject=subject,
        academic_level=academic_level,
        taxonomy=taxonomy,
    )
    if previous:
        prior_questions = {
            str(item.get("question_id")): item
            for item in previous.get("questions", [])
            if isinstance(item, Mapping) and item.get("question_id")
        }
        for question in analysis["questions"]:
            prior = prior_questions.get(str(question.get("question_id")))
            if not isinstance(prior, Mapping):
                continue
            override = prior.get("teacher_override")
            if isinstance(override, Mapping):
                question["teacher_override"] = dict(override)
                question["review_status"] = str(
                    override.get("review_status", prior.get("review_status", "teacher_reviewed"))
                )
    if consent_to_cloud:
        _enrich(analysis, client=client)
    _build_question_crops(store, record, analysis)
    _write_analysis_files(store, record, analysis)
    _append_audit(
        store,
        record,
        {
            "event": "analysis_created",
            "cloud_consent": consent_to_cloud,
            "model": analysis["model"].get("model"),
        },
    )
    store.set_analysis_summary(
        record.job_id,
        status=str(analysis["model"]["status"]),
        question_count=len(analysis["questions"]),
    )
    return analysis


def generate_practice(store: HistoryStore, record: HistoryRecord) -> dict[str, Any]:
    """Create local, evidence-linked practice starters without a cloud request."""

    analysis = read_analysis(store, record, create_preview=False)
    existing = analysis.get("practice_items")
    if isinstance(existing, Sequence) and not isinstance(existing, (str, bytes)) and existing:
        return read_practice(store, record)
    items: list[dict[str, Any]] = []
    for question in analysis.get("questions", [])[:20]:
        if not isinstance(question, Mapping):
            continue
        identifier = str(question.get("question_id", ""))
        text = " ".join(str(question.get("text", "")).split())[:900]
        if not _QUESTION_ID.fullmatch(identifier) or not text:
            continue
        items.append(
            {
                "source_question_id": identifier,
                "prompt": (
                    f"Teacher review required: draft one parallel practice question for "
                    f"{question.get('display_number', identifier)} using the same assessed skill. "
                    f"Base it only on this OCR evidence: {text}"
                ),
                "review_status": "needs_teacher_review",
                "source": "local_deterministic_starter",
            }
        )
    analysis["practice_items"] = items
    _write_analysis_files(store, record, analysis)
    _append_audit(store, record, {"event": "local_practice_generated", "count": len(items)})
    return read_practice(store, record)


def update_question(
    store: HistoryStore,
    record: HistoryRecord,
    question_id: str,
    updates: Mapping[str, Any],
) -> dict[str, Any]:
    if not _QUESTION_ID.fullmatch(question_id):
        raise EducationError("Question record not found")
    analysis = read_analysis(store, record, create_preview=False)
    allowed = {
        "topic",
        "subtopic",
        "difficulty",
        "difficulty_reason",
        "cognitive_skill",
        "marks",
        "review_status",
    }
    invalid = set(updates) - allowed
    if invalid:
        raise EducationError("Unsupported teacher review field")
    target = next(
        (item for item in analysis["questions"] if item.get("question_id") == question_id),
        None,
    )
    if not isinstance(target, dict):
        raise EducationError("Question record not found")
    override = target.setdefault("teacher_override", {})
    if not isinstance(override, dict):
        override = target["teacher_override"] = {}
    changed: dict[str, Any] = {}
    for name, value in updates.items():
        if name == "difficulty" and value not in _DIFFICULTIES:
            raise EducationError("Invalid difficulty")
        if name == "review_status" and value not in _REVIEW_STATUSES:
            raise EducationError("Invalid teacher review status")
        if name == "marks" and (not isinstance(value, int) or not 0 <= value <= 100):
            raise EducationError("Marks must be a whole number between 0 and 100")
        if name not in {"marks", "review_status", "difficulty"} and value is not None:
            value = " ".join(str(value).split())
            if len(value) > 500:
                raise EducationError("A teacher review value is too long")
        override[name] = value
        changed[name] = value
    target["review_status"] = str(override.get("review_status", "teacher_reviewed"))
    analysis["analysis"] = _aggregate_questions(analysis["questions"])
    _write_analysis_files(store, record, analysis)
    _append_audit(store, record, {"event": "teacher_override", "question_id": question_id, "changes": changed})
    return analysis


def read_practice(store: HistoryStore, record: HistoryRecord) -> dict[str, Any]:
    """Return only already validated, teacher-reviewable practice suggestions.

    Practice is produced during the explicitly consented analysis operation. This
    endpoint deliberately never triggers a silent cloud request.
    """

    analysis = read_analysis(store, record, create_preview=False)
    items = analysis.get("practice_items")
    return {
        "document_id": record.job_id,
        "model": analysis.get("model", {}).get("model"),
        "consent_to_cloud": bool(analysis.get("model", {}).get("consent_to_cloud")),
        "items": list(items) if isinstance(items, Sequence) and not isinstance(items, (str, bytes)) else [],
        "review_required": True,
    }


def aggregate_history(analyses: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    topic_sources: defaultdict[str, list[str]] = defaultdict(list)
    for analysis in analyses:
        for question in analysis.get("questions", []):
            if not isinstance(question, Mapping):
                continue
            topic = _effective_value(question, "topic")
            if not topic:
                continue
            topic_sources[str(topic)].append(str(question.get("question_id")))
    return {
        "recurring_topics": [
            {"topic": topic, "occurrences": len(question_ids), "source_question_ids": question_ids}
            for topic, question_ids in sorted(topic_sources.items(), key=lambda item: (-len(item[1]), item[0]))
            if len(question_ids) > 1
        ]
    }


def analysis_download_path(store: HistoryStore, record: HistoryRecord, format_name: str) -> Path:
    filenames = {"json": "analysis.json", "markdown": "analysis.md", "docx": "analysis.docx"}
    filename = filenames.get(format_name)
    if filename is None:
        raise EducationError("Unsupported analysis download format")
    path = _analysis_root(store, record) / filename
    if not path.is_file():
        raise EducationError("Create Past Paper Intelligence before downloading this analysis")
    return path


def source_preview_path(store: HistoryStore, record: HistoryRecord, page: int) -> Path:
    if page < 1 or page > 2000:
        raise EducationError("Invalid source page")
    manifest = store.read_manifest(record)
    if not manifest or not isinstance(manifest.get("archive_root"), str):
        raise EducationError("Source previews are unavailable for this History record")
    archive_root = store._resolve_history_path(manifest["archive_root"])  # noqa: SLF001
    path = archive_root / "source-pages" / f"page-{page:03d}.png"
    if not path.is_file():
        raise EducationError("Source preview is unavailable; reprocess the document to retain it")
    return path


def question_crop_path(store: HistoryStore, record: HistoryRecord, question_id: str) -> Path:
    if not _QUESTION_ID.fullmatch(question_id):
        raise EducationError("Question record not found")
    path = _analysis_root(store, record) / "crops" / f"{question_id}.png"
    if not path.is_file():
        raise EducationError("Question crop is unavailable")
    return path

"""Deterministic teaching services over immutable OCR evidence."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import shutil
import zipfile
from collections import Counter
from html import escape
from pathlib import Path
from typing import Any, Mapping, Sequence

from docx import Document

from app.education.questions import segment_questions
from app.education.service import EducationError, read_analysis
from app.storage.history import HistoryNotFoundError, HistoryResultError, HistoryStore
from app.teaching.languages import language_reliability
from app.teaching.store import TeachingStore, result_hash

_SPACE = re.compile(r"\s+")
_TOKEN = re.compile(r"[^\W_]+", re.UNICODE)
_NUMBER = re.compile(r"\b\d+(?:[.,]\d+)?\b")


def _lines(result: Mapping[str, Any]) -> list[dict[str, Any]]:
    raw = result.get("lines")
    if not isinstance(raw, Sequence) or isinstance(raw, (str, bytes)):
        return []
    return [dict(value) for value in raw if isinstance(value, Mapping)]


def _confidence(line: Mapping[str, Any]) -> float | None:
    for key in ("retry_confidence", "confidence", "score"):
        try:
            value = float(line.get(key))
        except (TypeError, ValueError):
            continue
        if 0 <= value <= 1:
            return round(value, 4)
    return None


def _bbox(line: Mapping[str, Any]) -> list[float] | None:
    raw = line.get("normalized_bbox")
    if not isinstance(raw, Sequence) or isinstance(raw, (str, bytes)) or len(raw) != 4:
        return None
    try:
        values = [float(value) for value in raw]
    except (TypeError, ValueError):
        return None
    left, top, right, bottom = values
    return [round(value, 5) for value in values] if 0 <= left <= right <= 1 and 0 <= top <= bottom <= 1 else None


def _band(confidence: float | None) -> str:
    if confidence is None:
        return "unavailable"
    if confidence >= 0.9:
        return "high"
    if confidence >= 0.7:
        return "review"
    return "low"


def effective_view(result: Mapping[str, Any], corrections: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    """Apply approved overlays without mutating the OCR result object."""

    accepted = {
        (str(item.get("target_kind")), str(item.get("target_id"))): item
        for item in corrections
        if item.get("status") == "approved" and not item.get("stale")
    }
    lines: list[dict[str, Any]] = []
    for index, raw in enumerate(_lines(result)):
        line = dict(raw)
        line_id = str(line.get("line_id") or f"line-{index + 1}")
        line["line_id"] = line_id
        overlay = accepted.get(("line", line_id)) or accepted.get(("block", line_id))
        original_text = str(line.get("text", ""))
        replacement = overlay.get("replacement") if overlay else None
        if isinstance(replacement, Mapping):
            replacement = replacement.get("text")
        line["raw_text"] = original_text
        line["text"] = str(replacement) if replacement is not None else original_text
        line["corrected"] = overlay is not None
        line["correction_id"] = overlay.get("correction_id") if overlay else None
        line["confidence"] = _confidence(line)
        line["confidence_band"] = _band(line["confidence"])
        line["normalized_bbox"] = _bbox(line)
        line["page_number"] = int(line.get("page_number", int(line.get("page_index", 0)) + 1))
        lines.append(line)
    reconstructed = "\n".join(line["text"] for line in lines if line["text"].strip())
    document_overlay = accepted.get(("reconstructed_text", "document"))
    if document_overlay:
        value = document_overlay.get("replacement")
        reconstructed = str(value.get("text") if isinstance(value, Mapping) else value)
    return {"lines": lines, "reconstructed_text": reconstructed, "approved_correction_count": len(accepted)}


def review_workspace(store: TeachingStore, job_id: str) -> dict[str, Any]:
    record = store.history.get(job_id)
    result = store.history.read_result(record)
    if not isinstance(result, Mapping):
        raise EducationError("The saved OCR evidence is not a structured document")
    correction_set = store.list_corrections(job_id)
    effective = effective_view(result, correction_set["items"])
    reliability = language_reliability(record.language)
    if reliability["question_segmentation"] == "supported":
        try:
            analysis = read_analysis(store.history, record, create_preview=True)
            questions = analysis.get("questions", [])
        except EducationError:
            questions = segment_questions(result)["questions"]
    else:
        questions = []
    approved = {
        (item["target_kind"], item["target_id"]): item
        for item in correction_set["items"]
        if item["status"] == "approved" and not item["stale"]
    }
    effective_questions = []
    for raw_question in questions:
        question = dict(raw_question)
        question_overlay = approved.get(("question", str(question.get("question_id"))))
        marks_overlay = approved.get(("marks", str(question.get("question_id"))))
        if question_overlay:
            replacement = question_overlay.get("replacement")
            question["raw_text"] = question.get("text")
            question["text"] = replacement.get("text") if isinstance(replacement, Mapping) else str(replacement)
            question["text_correction_id"] = question_overlay["correction_id"]
        if marks_overlay:
            replacement = marks_overlay.get("replacement")
            corrected_marks = replacement.get("marks") if isinstance(replacement, Mapping) else replacement
            if isinstance(corrected_marks, int) and 0 <= corrected_marks <= 100:
                question["raw_marks"] = question.get("marks")
                question["marks"] = corrected_marks
                question["marks_status"] = "teacher_approved_overlay"
                question["marks_correction_id"] = marks_overlay["correction_id"]
        effective_questions.append(question)
    with store.connect() as connection:
        mapped_questions = {
            str(row[0])
            for row in connection.execute(
                "SELECT DISTINCT question_id FROM syllabus_mappings "
                "WHERE job_id = ? AND status = 'approved'",
                (job_id,),
            ).fetchall()
        }
    question_lines: dict[str, list[Mapping[str, Any]]] = {}
    for question in effective_questions:
        for line_id in question.get("source_line_ids", []):
            question_lines.setdefault(str(line_id), []).append(question)
    for line in effective["lines"]:
        evidence_types = {"text"}
        linked_questions = question_lines.get(str(line["line_id"]), [])
        if linked_questions:
            evidence_types.add("question")
        if any(question.get("marks") is not None for question in linked_questions):
            evidence_types.add("marks")
        if any(
            str(question.get("question_id")) not in mapped_questions
            for question in linked_questions
        ):
            evidence_types.add("unmapped")
        raw_type = str(
            line.get("block_type")
            or line.get("category")
            or line.get("type")
            or ""
        ).casefold()
        if "table" in raw_type:
            evidence_types.add("table")
        if any(value in raw_type for value in ("figure", "image", "chart", "diagram")):
            evidence_types.add("figure")
        line["evidence_types"] = sorted(evidence_types)
    page_numbers = sorted({int(line["page_number"]) for line in effective["lines"]}) or [1]
    issues = []
    for line in effective["lines"]:
        confidence = line["confidence"]
        if confidence is None or confidence < 0.9 or line["corrected"]:
            issues.append(
                {
                    "issue_id": f"issue-{line['line_id']}",
                    "line_id": line["line_id"],
                    "page": line["page_number"],
                    "text": line["text"],
                    "confidence": confidence,
                    "confidence_band": line["confidence_band"],
                    "correction_status": "approved" if line["corrected"] else "uncorrected",
                    "geometry_available": line["normalized_bbox"] is not None,
                    "evidence_types": line["evidence_types"],
                }
            )
    return {
        "version": 1,
        "job_id": job_id,
        "input_name": record.input_name,
        "source_hash": correction_set["source_hash"],
        "language": record.language or "en",
        "language_reliability": reliability,
        "pages": [
            {
                "page": page,
                "source_preview_url": f"/api/v1/history/{job_id}/source-pages/{page}",
                "lines": [line for line in effective["lines"] if line["page_number"] == page],
            }
            for page in page_numbers
        ],
        "questions": effective_questions,
        "reconstructed_text": effective["reconstructed_text"],
        "approved_correction_count": effective["approved_correction_count"],
        "corrections": correction_set["items"],
        "screen_reader_issues": issues,
        "confidence_legend": [
            {"band": "high", "label": "High confidence, 90 percent or above"},
            {"band": "review", "label": "Review recommended, 70 to 89 percent"},
            {"band": "low", "label": "Low confidence, below 70 percent"},
            {"band": "unavailable", "label": "Confidence unavailable"},
        ],
    }


def _extract_docx(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(data)
    parts: list[str] = []
    try:
        for index in range(len(document)):
            page = document[index]
            text_page = page.get_textpage()
            try:
                parts.append(text_page.get_text_range())
            finally:
                text_page.close()
                page.close()
    finally:
        document.close()
    return "\n".join(parts)


def parse_syllabus(data: bytes, filename: str, content_type: str | None = None) -> dict[str, Any]:
    suffix = Path(filename).suffix.casefold()
    checksum = hashlib.sha256(data).hexdigest()
    warnings: list[str] = []
    inferred = suffix in {".pdf", ".docx"}
    if suffix == ".csv":
        rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        objectives = [
            {
                "objective_id": row.get("objective_id") or row.get("code") or f"objective-{index + 1}",
                "parent_id": row.get("parent_id") or None,
                "code": row.get("code"),
                "label": row.get("label") or row.get("description") or "",
                "description": row.get("description") or "",
                "aliases": [part.strip() for part in (row.get("aliases") or "").split("|") if part.strip()],
            }
            for index, row in enumerate(rows)
        ]
        source_type = "csv"
    elif suffix in {".json", ".case"}:
        value = json.loads(data.decode("utf-8-sig"))
        if isinstance(value, Mapping) and isinstance(value.get("CFItems"), Sequence):
            objectives = [
                {
                    "objective_id": item.get("identifier") or item.get("uri", {}).get("identifier"),
                    "parent_id": None,
                    "code": item.get("fullStatement"),
                    "label": item.get("humanCodingScheme") or item.get("fullStatement"),
                    "description": item.get("fullStatement", ""),
                    "aliases": [],
                }
                for item in value["CFItems"]
                if isinstance(item, Mapping)
            ]
            source_type = "ims_case_json"
        else:
            raw = value.get("objectives", []) if isinstance(value, Mapping) else value
            objectives = list(raw) if isinstance(raw, Sequence) else []
            source_type = "json"
    elif suffix in {".pdf", ".docx"}:
        text = _extract_pdf(data) if suffix == ".pdf" else _extract_docx(data)
        candidates = [line.strip(" \t-•") for line in text.splitlines() if len(line.strip()) >= 8]
        objectives = [
            {
                "objective_id": f"inferred-{index + 1}",
                "parent_id": None,
                "code": None,
                "label": line[:256],
                "description": line,
                "aliases": [],
                "inferred": True,
                "confirmed": False,
            }
            for index, line in enumerate(candidates[:500])
        ]
        warnings.append("Objectives were inferred from document text and require teacher confirmation before publication.")
        source_type = suffix[1:]
    else:
        raise ValueError("Supported syllabus formats are CSV, JSON, IMS CASE JSON, PDF, and DOCX")
    if not objectives:
        raise ValueError("No syllabus objectives could be extracted")
    for item in objectives:
        item["inferred"] = bool(item.get("inferred", inferred))
        item["confirmed"] = bool(item.get("confirmed", not inferred))
    return {"source_type": source_type, "source_checksum": checksum, "objectives": objectives, "warnings": warnings, "status": "draft"}


def parse_taxonomy(data: bytes, filename: str) -> dict[str, Any]:
    suffix = Path(filename).suffix.casefold()
    if suffix == ".csv":
        rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        nodes = [
            {
                "node_id": row.get("node_id") or f"topic-{index + 1}",
                "parent_id": row.get("parent_id") or None,
                "label": row.get("label") or "",
                "aliases": [part.strip() for part in (row.get("aliases") or "").split("|") if part.strip()],
                "description": row.get("description") or "",
                "sort_order": int(row.get("sort_order") or index),
            }
            for index, row in enumerate(rows)
        ]
    elif suffix == ".json":
        value = json.loads(data.decode("utf-8-sig"))
        nodes = value.get("nodes", []) if isinstance(value, Mapping) else value
        if not isinstance(nodes, Sequence) or isinstance(nodes, (str, bytes)):
            raise ValueError("Taxonomy JSON must contain a nodes array")
        nodes = list(nodes)
    else:
        raise ValueError("Taxonomy import supports CSV and JSON")
    if not nodes:
        raise ValueError("The taxonomy import contains no topic nodes")
    return {"nodes": nodes, "source_hash": hashlib.sha256(data).hexdigest()}


def taxonomy_export(taxonomy: Mapping[str, Any], format_name: str) -> tuple[bytes, str, str]:
    if format_name == "json":
        return (
            json.dumps(taxonomy, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8"),
            "application/json",
            "taxonomy.json",
        )
    if format_name != "csv":
        raise ValueError("Taxonomy export format is invalid")
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=["node_id", "parent_id", "label", "aliases", "description", "sort_order"])
    writer.writeheader()
    for node in taxonomy.get("nodes", []):
        writer.writerow(
            {
                "node_id": node.get("node_id"),
                "parent_id": node.get("parent_id") or "",
                "label": node.get("label"),
                "aliases": "|".join(node.get("aliases", [])),
                "description": node.get("description") or "",
                "sort_order": node.get("sort_order", 0),
            }
        )
    return output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", "taxonomy.csv"


def normalize_question(text: str, language: str = "en") -> str:
    folded = _SPACE.sub(" ", text.casefold()).strip()
    return " ".join(_TOKEN.findall(folded))


def _shingles(text: str, width: int = 3) -> set[tuple[str, ...]]:
    tokens = normalize_question(text).split()
    if len(tokens) < width:
        return {tuple(tokens)} if tokens else set()
    return {tuple(tokens[index : index + width]) for index in range(len(tokens) - width + 1)}


def duplicate_matches(questions: Sequence[Mapping[str, Any]], *, threshold: float = 0.55, language: str = "en") -> list[dict[str, Any]]:
    matches: list[dict[str, Any]] = []
    for left_index, left in enumerate(questions):
        left_text = str(left.get("text", ""))
        left_normal = normalize_question(left_text, language)
        left_shingles = _shingles(left_text)
        for right in questions[left_index + 1 :]:
            right_text = str(right.get("text", ""))
            right_normal = normalize_question(right_text, language)
            right_shingles = _shingles(right_text)
            exact = bool(left_normal and left_normal == right_normal)
            union = left_shingles | right_shingles
            similarity = 1.0 if exact else (len(left_shingles & right_shingles) / len(union) if union else 0.0)
            if similarity < threshold:
                continue
            shared = [" ".join(value) for value in sorted(left_shingles & right_shingles)[:10]]
            left_numbers, right_numbers = _NUMBER.findall(left_text), _NUMBER.findall(right_text)
            matches.append(
                {
                    "left_question_key": str(left.get("question_key") or left.get("question_id")),
                    "right_question_key": str(right.get("question_key") or right.get("question_id")),
                    "score": round(similarity, 4),
                    "method": "exact_fingerprint" if exact else "token_shingle",
                    "shared_phrases": shared,
                    "changed_numbers": left_numbers != right_numbers,
                    "left_numbers": left_numbers,
                    "right_numbers": right_numbers,
                    "teacher_disposition": None,
                    "automatic_action": "none",
                }
            )
    return sorted(matches, key=lambda item: item["score"], reverse=True)


def _question_from_item(item: Mapping[str, Any]) -> dict[str, Any]:
    payload = item.get("payload") if isinstance(item.get("payload"), Mapping) else {}
    return {"item_id": item.get("item_id"), "question_id": item.get("question_id"), **payload}


def export_question_bank(store: TeachingStore, *, formats: Sequence[str], include_source_images: bool = False) -> dict[str, Any]:
    items = [_question_from_item(item) for item in store.list_bank(status="approved") if not item.get("stale")]
    if not items:
        raise ValueError("No current approved question-bank items are available")
    allowed = {"qti", "csv", "jsonl", "markdown", "docx"}
    requested = [value for value in dict.fromkeys(formats) if value in allowed]
    if not requested:
        raise ValueError("Select at least one supported export format")
    export_id = hashlib.sha256(json.dumps(items, sort_keys=True).encode()).hexdigest()[:16]
    root = store.settings.history_dir / "question-bank-exports" / export_id
    root.mkdir(parents=True, exist_ok=True)
    files: list[Path] = []
    included_source_images: list[str] = []
    if include_source_images:
        from app.education.service import source_preview_path

        images_root = root / "source-pages"
        images_root.mkdir(exist_ok=True)
        for item in items:
            source_job = item.get("source_job_id")
            source_page = item.get("source_page")
            if not isinstance(source_job, str) or not isinstance(source_page, int):
                continue
            try:
                record = store.history.get(source_job)
                source = source_preview_path(store.history, record, source_page)
            except (EducationError, HistoryNotFoundError, HistoryResultError, OSError):
                continue
            destination = images_root / f"{source_job[:8]}-page-{source_page:03d}.png"
            if not destination.exists():
                shutil.copy2(source, destination)
            files.append(destination)
            included_source_images.append(destination.relative_to(root).as_posix())
    if "csv" in requested:
        path = root / "question-bank.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["item_id", "question_id", "text", "marks", "topic", "difficulty", "language", "source_job_id", "source_page"])
            writer.writeheader()
            for item in items:
                writer.writerow({name: item.get(name) for name in writer.fieldnames})
        files.append(path)
    if "jsonl" in requested:
        path = root / "question-bank.jsonl"
        path.write_text("\n".join(json.dumps(item, ensure_ascii=False, sort_keys=True) for item in items) + "\n", encoding="utf-8")
        files.append(path)
    if "markdown" in requested:
        path = root / "question-bank.md"
        chunks = ["# PredixaLearn teacher-approved question bank", "", "> Historical teaching evidence only. This material does not predict future examinations.", ""]
        for index, item in enumerate(items, 1):
            chunks.extend([f"## {index}. {item.get('text', 'Untitled question')}", "", f"- Marks: {item.get('marks', 'Not specified')}", f"- Topic: {item.get('topic', 'Unmapped')}", f"- Source: {item.get('source_job_id', 'Local History')} / page {item.get('source_page', '?')}", ""])
        path.write_text("\n".join(chunks), encoding="utf-8")
        files.append(path)
    if "docx" in requested:
        path = root / "question-bank.docx"
        document = Document()
        document.add_heading("PredixaLearn teacher-approved question bank", 0)
        document.add_paragraph("Historical teaching evidence only. This material does not predict future examinations.")
        for index, item in enumerate(items, 1):
            document.add_heading(f"{index}. {item.get('text', 'Untitled question')}", level=1)
            document.add_paragraph(f"Marks: {item.get('marks', 'Not specified')} | Topic: {item.get('topic', 'Unmapped')}")
        document.save(path)
        files.append(path)
    if "qti" in requested:
        path = root / "question-bank-qti-2.1.zip"
        resources = []
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for index, item in enumerate(items, 1):
                identifier = f"predixalearn-{item.get('item_id', index)}"
                filename = f"items/{identifier}.xml"
                body = escape(str(item.get("text", "")))
                xml = f'''<?xml version="1.0" encoding="UTF-8"?>\n<assessmentItem xmlns="http://www.imsglobal.org/xsd/imsqti_v2p1" identifier="{identifier}" title="Question {index}" adaptive="false" timeDependent="false"><itemBody><p>{body}</p></itemBody></assessmentItem>'''
                archive.writestr(filename, xml)
                resources.append((identifier, filename))
            manifest_resources = "".join(f'<resource identifier="res-{identifier}" type="imsqti_item_xmlv2p1" href="{filename}"><file href="{filename}"/></resource>' for identifier, filename in resources)
            archive.writestr("imsmanifest.xml", f'''<?xml version="1.0" encoding="UTF-8"?><manifest xmlns="http://www.imsglobal.org/xsd/imscp_v1p1" identifier="predixalearn-question-bank"><resources>{manifest_resources}</resources></manifest>''')
        files.append(path)
    checksums = {path.relative_to(root).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest() for path in files}
    manifest = {"version": 1, "export_id": export_id, "item_count": len(items), "source_images_requested": bool(include_source_images), "source_images_included": included_source_images, "files": checksums}
    manifest_path = root / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    archive_path = root / "predixalearn-question-bank.zip"
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in [*files, manifest_path]:
            archive.write(path, path.relative_to(root).as_posix())
    return {**manifest, "download_path": str(archive_path), "download_url": f"/api/v1/question-bank/exports/{export_id}"}


def revision_pack_export_path(
    store: TeachingStore,
    pack_id: str,
    format_name: str,
) -> Path:
    """Build a student-safe standalone revision export from current approved items."""

    if format_name not in {"markdown", "json", "docx"}:
        raise ValueError("Revision-pack export format is invalid")
    pack = store.get_revision_pack(pack_id)
    questions: list[dict[str, Any]] = []
    for item_id in pack["source_item_ids"]:
        try:
            item = store.get_bank_item(str(item_id))
        except KeyError:
            continue
        if item["status"] != "approved" or item.get("stale"):
            continue
        payload = item["payload"] if isinstance(item.get("payload"), Mapping) else {}
        questions.append(
            {
                key: payload.get(key)
                for key in (
                    "question_id",
                    "text",
                    "topic",
                    "objective",
                    "difficulty",
                    "language",
                    "provenance",
                    "accessibility_description",
                )
                if payload.get(key) is not None
            }
        )
    if pack["status"] == "approved" and len(questions) != len(pack["source_item_ids"]):
        raise ValueError(
            "The approved pack references a missing, stale, or unapproved question-bank item"
        )
    root = store.settings.history_dir / "revision-pack-exports" / pack_id
    root.mkdir(parents=True, exist_ok=True)
    safe_payload = {
        "version": 1,
        "pack_id": pack_id,
        "title": pack["title"],
        "status": pack["status"],
        "language": pack["language"],
        "introduction": pack["payload"].get("introduction"),
        "questions": questions,
        "statement": "Teacher-reviewed historical practice; not a future-exam prediction.",
        "student_safe": True,
    }
    if format_name == "json":
        path = root / "revision-pack.json"
        path.write_text(
            json.dumps(safe_payload, ensure_ascii=False, indent=2, sort_keys=True),
            encoding="utf-8",
        )
        return path
    if format_name == "markdown":
        path = root / "revision-pack.md"
        lines = [
            f"# {pack['title']}",
            "",
            str(safe_payload["introduction"] or ""),
            "",
            f"> {safe_payload['statement']}",
            "",
        ]
        for index, question in enumerate(questions, 1):
            lines.extend(
                [
                    f"## Practice question {index}",
                    "",
                    str(question.get("text", "Question text unavailable")),
                    "",
                    f"- Topic: {question.get('topic', 'Unmapped')}",
                    f"- Objective: {question.get('objective', 'Unmapped')}",
                    f"- Difficulty: {question.get('difficulty', 'Unclassified')}",
                    "",
                ]
            )
        path.write_text("\n".join(lines), encoding="utf-8")
        return path
    path = root / "revision-pack.docx"
    document = Document()
    document.add_heading(pack["title"], 0)
    if safe_payload["introduction"]:
        document.add_paragraph(str(safe_payload["introduction"]))
    document.add_paragraph(str(safe_payload["statement"]))
    for index, question in enumerate(questions, 1):
        document.add_heading(f"Practice question {index}", level=1)
        document.add_paragraph(str(question.get("text", "Question text unavailable")))
        document.add_paragraph(
            f"Topic: {question.get('topic', 'Unmapped')} | "
            f"Objective: {question.get('objective', 'Unmapped')} | "
            f"Difficulty: {question.get('difficulty', 'Unclassified')}"
        )
    document.save(path)
    return path


def compare_papers(history: HistoryStore, teaching: TeachingStore, job_ids: Sequence[str]) -> dict[str, Any]:
    unique = list(dict.fromkeys(job_ids))
    if not 2 <= len(unique) <= 50:
        raise ValueError("Select between 2 and 50 papers")
    papers, topic_counts, difficulty_counts, warnings = [], Counter(), Counter(), []
    cognitive_counts: Counter[str] = Counter()
    question_type_counts: Counter[str] = Counter()
    all_questions: list[dict[str, Any]] = []
    source_hashes: dict[str, str] = {}
    for job_id in unique:
        record = history.get(job_id)
        result = history.read_result(record)
        source_hashes[job_id] = result_hash(result)
        try:
            analysis = read_analysis(history, record, create_preview=False)
            questions = analysis.get("questions", [])
        except EducationError:
            questions = segment_questions(result)["questions"] if isinstance(result, Mapping) else []
            warnings.append(f"{record.input_name} has deterministic segmentation only; analyze it for topic comparisons.")
        marks_total, low_confidence, unmapped = 0, 0, 0
        for question in questions:
            teacher = question.get("teacher_override") if isinstance(question.get("teacher_override"), Mapping) else {}
            ai = question.get("ai") if isinstance(question.get("ai"), Mapping) else {}
            topic = teacher.get("topic", ai.get("topic")) or "Unmapped"
            difficulty = teacher.get("difficulty", ai.get("difficulty")) or "unclassified"
            cognitive = teacher.get("cognitive_skill", ai.get("cognitive_skill")) or "unclassified"
            question_type = teacher.get("question_type", ai.get("question_type")) or "unclassified"
            topic_counts[str(topic)] += 1
            difficulty_counts[str(difficulty)] += 1
            cognitive_counts[str(cognitive)] += 1
            question_type_counts[str(question_type)] += 1
            unmapped += int(topic == "Unmapped")
            marks = teacher.get("marks", question.get("marks"))
            marks_total += marks if isinstance(marks, int) else 0
            confidence = question.get("ocr_confidence")
            low_confidence += int(isinstance(confidence, (int, float)) and confidence < 0.7)
            all_questions.append({"question_key": f"{job_id}:{question.get('question_id')}", "job_id": job_id, "paper": record.input_name, **question})
        with teaching.connect() as connection:
            correction_workload = int(
                connection.execute(
                    "SELECT COUNT(*) FROM correction_sets WHERE job_id = ? "
                    "AND status IN ('draft', 'reviewed')",
                    (job_id,),
                ).fetchone()[0]
            )
            approved_mappings = int(
                connection.execute(
                    "SELECT COUNT(*) FROM syllabus_mappings WHERE job_id = ? AND status = 'approved'",
                    (job_id,),
                ).fetchone()[0]
            )
        papers.append({"job_id": job_id, "input_name": record.input_name, "question_count": len(questions), "marks_total": marks_total, "low_confidence": low_confidence, "unmapped": unmapped, "approved_syllabus_mappings": approved_mappings, "pending_corrections": correction_workload, "quality_score": record.quality_score, "warning_count": record.warning_count})
    duplicates = duplicate_matches(all_questions)
    return {
        "paper_count": len(papers),
        "papers": papers,
        "source_hashes": source_hashes,
        "topics": [{"topic": key, "question_count": value} for key, value in topic_counts.most_common()],
        "difficulty": [{"difficulty": key, "question_count": value} for key, value in difficulty_counts.most_common()],
        "cognitive_skills": [{"skill": key, "question_count": value} for key, value in cognitive_counts.most_common()],
        "question_types": [{"question_type": key, "question_count": value} for key, value in question_type_counts.most_common()],
        "duplicates": duplicates,
        "warnings": warnings,
        "statement": "This report summarizes historical evidence and does not predict future examination questions.",
    }


def corrected_export_path(store: TeachingStore, job_id: str, format_name: str) -> Path:
    if format_name not in {"markdown", "json", "docx"}:
        raise ValueError("Corrected export format is invalid")
    workspace = review_workspace(store, job_id)
    root = store.settings.history_dir / "jobs" / job_id / "teaching-exports"
    root.mkdir(parents=True, exist_ok=True)
    if format_name == "json":
        path = root / "teacher-approved-effective-view.json"
        payload = {
            "version": 1,
            "job_id": job_id,
            "ocr_result_sha256": workspace["source_hash"],
            "raw_ocr_mutated": False,
            "approved_correction_count": workspace["approved_correction_count"],
            "language": workspace["language"],
            "questions": workspace["questions"],
            "lines": [line for page in workspace["pages"] for line in page["lines"]],
            "reconstructed_text": workspace["reconstructed_text"],
        }
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
        return path
    if format_name == "markdown":
        path = root / "teacher-approved-effective-view.md"
        parts = [
            "# Teacher-approved effective reconstruction",
            "",
            f"> Immutable OCR evidence SHA-256: `{workspace['source_hash']}`. Corrections are approved overlays; raw OCR was not changed.",
            "",
            workspace["reconstructed_text"],
        ]
        path.write_text("\n".join(parts), encoding="utf-8")
        return path
    path = root / "teacher-approved-effective-view.docx"
    document = Document()
    document.add_heading("Teacher-approved effective reconstruction", 0)
    document.add_paragraph(
        f"Immutable OCR evidence SHA-256: {workspace['source_hash']}. "
        "Corrections are approved overlays; raw OCR was not changed."
    )
    for paragraph in workspace["reconstructed_text"].splitlines():
        if paragraph.strip():
            document.add_paragraph(paragraph)
    document.save(path)
    return path

"""Structured Markdown and deterministic DOCX construction for OCR results."""

from __future__ import annotations

import html
import importlib.resources
import json
import math
import os
import posixpath
import re
import subprocess
import uuid
import zipfile
from collections import defaultdict
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from app.core.capabilities import validate_document_profile
from app.core.config import Settings, get_settings
from app.core.geometry import (
    bbox_iou as geometry_bbox_iou,
)
from app.core.geometry import (
    bbox_values as geometry_bbox_values,
)
from app.core.geometry import (
    overlap_ratio as geometry_overlap_ratio,
)
from app.core.geometry import (
    preferred_bbox,
)
from app.documents.docx import pandoc_executable
from app.documents.naming import safe_output_stem
from app.documents.safety import escape_markdown_text

_URL_OR_PROTECTED = re.compile(
    r"(https?://\S+|www\.\S+|`[^`]*`|\$\$.*?\$\$|\$[^$\n]+\$|<[^>]+>)",
    re.IGNORECASE | re.DOTALL,
)
_WORD = re.compile(r"\b[a-z][a-z'-]{3,}\b")
_MATH_FUNCTIONS = frozenset(
    {"sqrt", "sin", "cos", "tan", "log", "ln", "exp", "lim", "sum", "prod", "int"}
)
_MATH_CONNECTOR = re.compile(
    r"\s+(?:at|where|when|for|and|or|but|with|from|on)\s+(?=[A-Za-z])",
    re.IGNORECASE,
)
_TEXT_SUPPLEMENTAL_TYPES = (
    "formula",
    "equation",
    "image",
    "figure",
    "chart",
    "illustration",
    "graphic",
)
_IMAGE_MARKDOWN = re.compile(r"!\[[^\]]*\]\(([^)]+)\)")
_PAGE_MARKER = re.compile(r"^Source page (\d+)$", re.IGNORECASE)
_RAW_IMAGE_TAG = re.compile(r"<img\b[^>]*>", re.IGNORECASE)
_RAW_SCRIPT_TAG = re.compile(r"<(script|style)\b.*?</\1\s*>", re.IGNORECASE | re.DOTALL)
_EVENT_ATTRIBUTE = re.compile(r"\s+on[a-z]+\s*=\s*(['\"]).*?\1", re.IGNORECASE | re.DOTALL)
_MATH_COMMAND = re.compile(r"(?<!\\)\\([A-Za-z]+)")
_MATH_ENVIRONMENT = re.compile(r"\\(begin|end)\s*\{([^{}]+)\}")
_DISPLAY_MATH_BLOCK = re.compile(
    r"^\$\$[ \t]*\r?\n(.*?)\r?\n\$\$[ \t]*$",
    re.MULTILINE | re.DOTALL,
)
_ALLOWED_MATH_ENVIRONMENTS = frozenset(
    {
        "aligned",
        "alignedat",
        "array",
        "matrix",
        "pmatrix",
        "bmatrix",
        "Bmatrix",
        "vmatrix",
        "Vmatrix",
        "cases",
        "gathered",
        "split",
    }
)
_FORBIDDEN_MATH_COMMANDS = frozenset(
    {
        "clearpage",
        "class",
        "def",
        "documentclass",
        "href",
        "htmlclass",
        "htmlid",
        "htmlstyle",
        "id",
        "include",
        "includegraphics",
        "input",
        "newcommand",
        "newpage",
        "openout",
        "pagebreak",
        "read",
        "renewcommand",
        "require",
        "style",
        "url",
        "usepackage",
        "write",
    }
)
_ALLOWED_MATH_COMMANDS = frozenset(
    {
        # Structure and sizing.
        "begin",
        "end",
        "left",
        "right",
        "middle",
        "big",
        "Big",
        "bigg",
        "Bigg",
        "bigl",
        "bigr",
        "Bigl",
        "Bigr",
        "biggl",
        "biggr",
        "Biggl",
        "Biggr",
        "displaystyle",
        "textstyle",
        "scriptstyle",
        "scriptscriptstyle",
        "substack",
        "hline",
        "cline",
        "tag",
        # Fractions, roots, annotations, and text.
        "frac",
        "dfrac",
        "tfrac",
        "cfrac",
        "sqrt",
        "binom",
        "dbinom",
        "tbinom",
        "overset",
        "underset",
        "stackrel",
        "overline",
        "underline",
        "overbrace",
        "underbrace",
        "widehat",
        "widetilde",
        "hat",
        "bar",
        "vec",
        "dot",
        "ddot",
        "breve",
        "check",
        "acute",
        "grave",
        "tilde",
        "text",
        "textrm",
        "textnormal",
        "textbf",
        "textit",
        "mathrm",
        "mathbf",
        "mathit",
        "mathsf",
        "mathtt",
        "mathcal",
        "mathbb",
        "mathfrak",
        "boldsymbol",
        "operatorname",
        # Large operators and common functions.
        "sum",
        "prod",
        "coprod",
        "int",
        "iint",
        "iiint",
        "iiiint",
        "oint",
        "lim",
        "limsup",
        "liminf",
        "min",
        "max",
        "inf",
        "sup",
        "arg",
        "det",
        "dim",
        "gcd",
        "hom",
        "ker",
        "Pr",
        "log",
        "ln",
        "lg",
        "exp",
        "sin",
        "cos",
        "tan",
        "cot",
        "sec",
        "csc",
        "arcsin",
        "arccos",
        "arctan",
        "sinh",
        "cosh",
        "tanh",
        "coth",
        # Greek letters.
        "alpha",
        "beta",
        "gamma",
        "delta",
        "epsilon",
        "varepsilon",
        "zeta",
        "eta",
        "theta",
        "vartheta",
        "iota",
        "kappa",
        "lambda",
        "mu",
        "nu",
        "xi",
        "omicron",
        "pi",
        "varpi",
        "rho",
        "varrho",
        "sigma",
        "varsigma",
        "tau",
        "upsilon",
        "phi",
        "varphi",
        "chi",
        "psi",
        "omega",
        "Gamma",
        "Delta",
        "Theta",
        "Lambda",
        "Xi",
        "Pi",
        "Sigma",
        "Upsilon",
        "Phi",
        "Psi",
        "Omega",
        # Arithmetic, relations, sets, and arrows.
        "cdot",
        "times",
        "div",
        "pm",
        "mp",
        "ast",
        "star",
        "circ",
        "bullet",
        "oplus",
        "ominus",
        "otimes",
        "oslash",
        "le",
        "leq",
        "ge",
        "geq",
        "ne",
        "neq",
        "approx",
        "sim",
        "simeq",
        "cong",
        "equiv",
        "propto",
        "in",
        "notin",
        "ni",
        "subset",
        "subseteq",
        "supset",
        "supseteq",
        "parallel",
        "perp",
        "mid",
        "nmid",
        "ll",
        "gg",
        "forall",
        "exists",
        "nexists",
        "neg",
        "land",
        "lor",
        "cap",
        "cup",
        "setminus",
        "emptyset",
        "varnothing",
        "to",
        "mapsto",
        "gets",
        "rightarrow",
        "leftarrow",
        "leftrightarrow",
        "Rightarrow",
        "Leftarrow",
        "Leftrightarrow",
        "uparrow",
        "downarrow",
        "updownarrow",
        "longrightarrow",
        "longleftarrow",
        "longleftrightarrow",
        # Delimiters, spacing, and common symbols.
        "langle",
        "rangle",
        "lfloor",
        "rfloor",
        "lceil",
        "rceil",
        "lvert",
        "rvert",
        "vert",
        "Vert",
        "quad",
        "qquad",
        "enspace",
        "thinspace",
        "ldots",
        "cdots",
        "vdots",
        "ddots",
        "dots",
        "infty",
        "partial",
        "nabla",
        "ell",
        "hbar",
        "imath",
        "jmath",
        "Re",
        "Im",
        "angle",
        "degree",
        "prime",
    }
)


def _bbox_values(raw: Any) -> tuple[float, float, float, float] | None:
    return geometry_bbox_values(preferred_bbox(raw))


def _reading_key(item: Mapping[str, Any]) -> tuple[int, float, float, int]:
    page = int(item.get("page_index", 0) or 0)
    bbox = _bbox_values(item)
    y = bbox[1] if bbox else float("inf")
    x = bbox[0] if bbox else float("inf")
    index = int(item.get("index", 0) or 0)
    return page, y, x, index


def _bbox_iou(first: Any, second: Any) -> float:
    return geometry_bbox_iou(first, second)


def suppress_adjacent_duplicates(
    elements: Iterable[Mapping[str, Any]],
    *,
    preserve_order: bool = False,
) -> list[dict[str, Any]]:
    """Remove only overlapping adjacent exact/prefix duplicates."""
    source = elements if preserve_order else sorted(elements, key=_reading_key)
    ordered = [dict(item) for item in source]
    kept: list[dict[str, Any]] = []
    for item in ordered:
        content = re.sub(r"\s+", " ", str(item.get("content", ""))).strip()
        if kept and content:
            prior = kept[-1]
            prior_content = re.sub(r"\s+", " ", str(prior.get("content", ""))).strip()
            same_page = int(prior.get("page_index", 0)) == int(item.get("page_index", 0))
            overlap = _bbox_iou(prior, item) >= 0.85
            folded = content.casefold()
            prior_folded = prior_content.casefold()
            exact = folded == prior_folded
            prefix = (
                min(len(folded), len(prior_folded)) >= 12
                and (
                    folded.startswith(prior_folded)
                    or prior_folded.startswith(folded)
                )
                and min(len(folded), len(prior_folded))
                / max(len(folded), len(prior_folded))
                >= 0.9
            )
            if same_page and overlap and (exact or prefix):
                if len(content) > len(prior_content):
                    kept[-1] = item
                continue
        kept.append(item)
    return kept


@dataclass(frozen=True, slots=True)
class Correction:
    original: str
    corrected: str
    page: int
    block: int
    kind: str
    confidence: float
    reason: list[str]

    def as_dict(self) -> dict[str, Any]:
        return {
            "original": self.original,
            "corrected": self.corrected,
            "page": self.page,
            "block": self.block,
            "kind": self.kind,
            "confidence": self.confidence,
            "reason": self.reason,
        }


class ConservativeCorrector:
    """Offline edit-distance-one correction with a deliberately high acceptance bar."""

    def __init__(self, enabled: bool = True) -> None:
        self.enabled = enabled
        self._symspell: Any | None = None
        self.warning: str | None = None
        self.review_suggestions: list[dict[str, Any]] = []

    def _load(self) -> Any | None:
        if not self.enabled:
            return None
        if self._symspell is not None:
            return self._symspell
        try:
            from symspellpy import SymSpell

            dictionary = importlib.resources.files("symspellpy").joinpath(
                "frequency_dictionary_en_82_765.txt"
            )
            symspell = SymSpell(max_dictionary_edit_distance=1, prefix_length=7)
            if not symspell.load_dictionary(str(dictionary), term_index=0, count_index=1):
                raise RuntimeError("SymSpell dictionary could not be loaded")
            bigrams = importlib.resources.files("symspellpy").joinpath(
                "frequency_bigramdictionary_en_243_342.txt"
            )
            if not symspell.load_bigram_dictionary(
                str(bigrams),
                term_index=0,
                count_index=2,
            ):
                raise RuntimeError("SymSpell bigram dictionary could not be loaded")
            self._symspell = symspell
            return symspell
        except (ImportError, OSError, RuntimeError) as exc:
            self.enabled = False
            self.warning = f"Automatic correction was skipped: {type(exc).__name__}"
            return None

    @staticmethod
    def _normalize_punctuation(text: str) -> tuple[str, bool]:
        changed = False
        lines: list[str] = []
        for line in text.splitlines() or [text]:
            updated = re.sub(r"[ \t]+([,;:!?])", r"\1", line)
            updated = re.sub(r"([,;:!?])(?=[A-Za-z])", r"\1 ", updated)
            updated = re.sub(r"(?<=[a-z])\.(?=[A-Z])", ". ", updated)
            updated = re.sub(r"([,;:])\1+", r"\1", updated)
            updated = re.sub(r"[ \t]{2,}", " ", updated)
            changed = changed or updated != line
            lines.append(updated)
        return "\n".join(lines), changed

    @staticmethod
    def _zipf(token: str) -> float:
        from wordfreq import zipf_frequency

        return float(zipf_frequency(token, "en", wordlist="best"))

    def _add_review(
        self,
        *,
        page: int,
        block: int,
        original: str,
        suggested: str | None,
        kind: str,
        confidence: float,
        reason: str,
    ) -> None:
        item = {
            "page": page,
            "block": block,
            "original": original,
            "suggested": suggested,
            "kind": kind,
            "confidence": round(confidence, 4),
            "reason": reason,
        }
        if item not in self.review_suggestions:
            self.review_suggestions.append(item)

    def _word_boundary_candidates(self, token: str, symspell: Any) -> list[tuple[float, list[str]]]:
        """Return the two strongest exact segmentations so ambiguity is measurable."""
        short_words = {"a", "i", "of", "to", "in", "on", "at", "by", "or", "as", "is", "it", "an", "be"}
        length = len(token)
        paths: dict[int, list[tuple[float, list[str]]]] = {0: [(0.0, [])]}
        for start in range(length):
            if start not in paths:
                continue
            for end in range(start + 1, min(length, start + 24) + 1):
                word = token[start:end].casefold()
                if word not in short_words and (len(word) < 3 or word not in symspell.words):
                    continue
                frequency = self._zipf(word)
                if word not in short_words and frequency < 3.0:
                    continue
                for score, words in paths[start]:
                    if len(words) >= 5:
                        continue
                    bigram_bonus = 0.0
                    if words:
                        count = int(symspell.bigrams.get(f"{words[-1]} {word}", 0))
                        phrase_frequency = self._zipf(f"{words[-1]} {word}")
                        if count <= 0 and phrase_frequency < 4.0:
                            continue
                        bigram_bonus = (
                            math.log1p(count) / 4.0 if count > 0 else phrase_frequency
                        )
                    paths.setdefault(end, []).append(
                        (score + frequency + bigram_bonus, [*words, word])
                    )
                    paths[end] = sorted(paths[end], reverse=True)[:4]
        return [item for item in sorted(paths.get(length, ()), reverse=True) if len(item[1]) >= 2][
            :2
        ]

    def _correct_segment(
        self,
        text: str,
        *,
        page: int,
        block: int,
    ) -> tuple[str, list[str], list[str], list[float]]:
        symspell = self._load()
        reasons: list[str] = []
        kinds: list[str] = []
        confidences: list[float] = []
        updated, punctuation_changed = self._normalize_punctuation(text)
        if punctuation_changed:
            reasons.append("deterministic punctuation/spacing normalization")
            kinds.append("punctuation")
            confidences.append(1.0)
        if symspell is None:
            return updated, reasons, kinds, confidences

        from symspellpy import Verbosity

        # Paddle may return equations as plain OCR rather than delimited LaTeX.
        # Protect the assignment/expression span while still allowing a typo in
        # surrounding prose to be corrected.
        formula_spans: list[tuple[int, int]] = []
        line_start = 0
        for line in updated.splitlines(keepends=True) or [updated]:
            body = line.rstrip("\r\n")
            equals = body.find("=")
            if equals >= 0:
                lhs = re.search(r"[A-Za-z][A-Za-z0-9_]*\s*$", body[:equals])
                start = lhs.start() if lhs else 0
                connector = _MATH_CONNECTOR.search(body, equals + 1)
                end = connector.start() if connector else len(body)
                formula_spans.append((line_start + start, line_start + end))
            line_start += len(line)

        def replace(match: re.Match[str]) -> str:
            token = match.group(0)
            if token.casefold() in _MATH_FUNCTIONS or any(
                start <= match.start() and match.end() <= end
                for start, end in formula_spans
            ):
                return token
            if token != token.lower() or "'" in token or "-" in token:
                return token
            source_frequency = self._zipf(token)
            if token in symspell.words or source_frequency >= 2.0:
                return token

            if len(token) >= 12 and token.isalpha() and source_frequency < 1.5:
                candidates = self._word_boundary_candidates(token, symspell)
                if candidates and (
                    len(candidates) == 1 or candidates[0][0] >= candidates[1][0] + 2.5
                ):
                    words = candidates[0][1]
                    replacement = " ".join(words)
                    reasons.append(f"high-confidence word-boundary restoration: {token} -> {replacement}")
                    kinds.append("word_boundary")
                    confidences.append(0.97)
                    return replacement
                if candidates:
                    words = candidates[0][1]
                    self._add_review(
                        page=page,
                        block=block,
                        original=token,
                        suggested=" ".join(words),
                        kind="word_boundary",
                        confidence=0.55,
                        reason="A possible word split was not strong enough to apply automatically",
                    )
            suggestions = symspell.lookup(
                token,
                Verbosity.ALL,
                max_edit_distance=1,
                include_unknown=False,
                transfer_casing=False,
            )
            candidates = [item for item in suggestions if item.distance == 1 and item.count >= 5000]
            if not candidates:
                return token
            candidates.sort(key=lambda item: (-item.count, item.term))
            best = candidates[0]
            if len(candidates) > 1 and best.count < candidates[1].count * 10:
                return token
            if best.term == token or not best.term.isalpha():
                return token
            candidate_frequency = self._zipf(best.term)
            left_words = _WORD.findall(updated[: match.start()])
            right_words = _WORD.findall(updated[match.end() :])
            left = left_words[-1].casefold() if left_words else None
            right = right_words[0].casefold() if right_words else None

            def context_score(candidate: str) -> float:
                score = 0.0
                if left:
                    score += math.log1p(symspell.bigrams.get(f"{left} {candidate}", 0))
                if right:
                    score += math.log1p(symspell.bigrams.get(f"{candidate} {right}", 0))
                return score

            source_context = context_score(token)
            candidate_context = context_score(best.term)
            frequency_gain = candidate_frequency - source_frequency
            context_improves = candidate_context > source_context + math.log(3)
            if candidate_frequency < 3.5 or frequency_gain < 1.5 or (
                (left or right) and source_context > 0 and not context_improves
            ):
                self._add_review(
                    page=page,
                    block=block,
                    original=token,
                    suggested=best.term,
                    kind="spelling",
                    confidence=0.6,
                    reason="A spelling candidate was contextually ambiguous",
                )
                return token
            reasons.append(f"unambiguous dictionary correction: {token} -> {best.term}")
            kinds.append("spelling")
            confidences.append(0.95 if context_improves else 0.9)
            return best.term

        return _WORD.sub(replace, updated), reasons, kinds, confidences

    def correct(self, text: str, *, page: int, block: int) -> tuple[str, Correction | None]:
        if not self.enabled or not text.strip():
            return text, None
        pieces = _URL_OR_PROTECTED.split(text)
        reasons: list[str] = []
        kinds: list[str] = []
        confidences: list[float] = []
        for index in range(0, len(pieces), 2):
            pieces[index], part_reasons, part_kinds, part_confidences = self._correct_segment(
                pieces[index],
                page=page,
                block=block,
            )
            reasons.extend(part_reasons)
            kinds.extend(part_kinds)
            confidences.extend(part_confidences)
        joined = "".join(pieces)
        corrected = joined.strip()
        if corrected != joined:
            reasons.append("leading/trailing whitespace normalization")
            kinds.append("whitespace")
            confidences.append(1.0)
        if corrected == text:
            self._add_grammar_reviews(corrected, page=page, block=block)
            return text, None
        self._add_grammar_reviews(corrected, page=page, block=block)
        unique_kinds = list(dict.fromkeys(kinds))
        return corrected, Correction(
            text,
            corrected,
            page,
            block,
            "+".join(unique_kinds) if unique_kinds else "normalization",
            round(min(confidences, default=1.0), 4),
            list(dict.fromkeys(reasons)),
        )

    def _add_grammar_reviews(self, text: str, *, page: int, block: int) -> None:
        for pattern, replacement, reason in (
            (r"\bshort not on\b", "short note on", "Possible noun-form error after 'short'"),
            (r"\b([A-Za-z]+)\s+\1\b", None, "Possible repeated word"),
        ):
            match = re.search(pattern, text, re.IGNORECASE)
            if not match:
                continue
            suggested = replacement
            if suggested is None:
                suggested = match.group(1)
            self._add_review(
                page=page,
                block=block,
                original=match.group(0),
                suggested=suggested,
                kind="grammar",
                confidence=0.65,
                reason=reason,
            )


def _escape_plain(text: str) -> str:
    """Compatibility façade for the shared OCR-to-Markdown safety primitive."""
    return escape_markdown_text(text)


def _safe_figure_alt(value: str, *, fallback: str, max_length: int = 180) -> str:
    """Return one bounded, Markdown-inert line for generated image syntax."""
    normalized = re.sub(r"\s+", " ", value).strip() or fallback
    if len(normalized) > max_length:
        normalized = normalized[: max_length - 1].rstrip() + "…"
    return _escape_plain(normalized)


def _is_escaped(value: str, index: int) -> bool:
    backslashes = 0
    position = index - 1
    while position >= 0 and value[position] == "\\":
        backslashes += 1
        position -= 1
    return backslashes % 2 == 1


def _strip_formula_delimiters(value: str) -> str:
    """Remove only matching wrappers so the caller can add one safe display wrapper."""
    formula = value.replace("\r\n", "\n").replace("\r", "\n").strip()
    wrappers = (("$$", "$$"), (r"\[", r"\]"), (r"\(", r"\)"), ("$", "$"))
    changed = True
    while changed and formula:
        changed = False
        for opening, closing in wrappers:
            if (
                formula.startswith(opening)
                and formula.endswith(closing)
                and len(formula) >= len(opening) + len(closing)
            ):
                formula = formula[len(opening) : len(formula) - len(closing)].strip()
                changed = True
                break
    return formula


def _formula_rejection_reason(formula: str) -> str | None:
    """Return a conservative rendering problem, or ``None`` for safe KaTeX input."""
    if not formula:
        return "the OCR formula is empty"
    if any(ord(character) < 32 and character not in "\n\t" for character in formula):
        return "the formula contains unsupported control characters"
    if any(character == "$" and not _is_escaped(formula, index) for index, character in enumerate(formula)):
        return "the formula contains an embedded math delimiter"
    for match in re.finditer(r"\\[()\[\]]", formula):
        if not _is_escaped(formula, match.start()):
            return "the formula contains an embedded math delimiter"

    bracket_stack: list[tuple[str, int]] = []
    closing_to_opening = {"}": "{", "]": "[", ")": "("}
    for index, character in enumerate(formula):
        if _is_escaped(formula, index):
            continue
        if character in "{[(":
            bracket_stack.append((character, index))
        elif character in "}])":
            if not bracket_stack or bracket_stack[-1][0] != closing_to_opening[character]:
                return f"the formula has an unmatched '{character}' delimiter"
            bracket_stack.pop()
    if bracket_stack:
        return f"the formula has an unmatched '{bracket_stack[-1][0]}' delimiter"

    environment_stack: list[str] = []
    for match in _MATH_ENVIRONMENT.finditer(formula):
        action, environment = match.groups()
        environment = environment.strip()
        if environment not in _ALLOWED_MATH_ENVIRONMENTS:
            return f"the formula uses unsupported environment '{environment}'"
        if action == "begin":
            environment_stack.append(environment)
        elif not environment_stack or environment_stack[-1] != environment:
            return f"the formula has a mismatched end for environment '{environment}'"
        else:
            environment_stack.pop()
    if environment_stack:
        return f"the formula does not close environment '{environment_stack[-1]}'"

    for command in _MATH_COMMAND.findall(formula):
        if command.casefold() in _FORBIDDEN_MATH_COMMANDS:
            return f"the formula uses unsafe command '\\{command}'"
        if command not in _ALLOWED_MATH_COMMANDS:
            return f"the formula uses unsupported command '\\{command}'"

    in_environment = bool(_MATH_ENVIRONMENT.search(formula))
    if r"\\" in formula and not in_environment:
        return "the formula uses a line break outside an approved math environment"
    if "&" in formula and not in_environment:
        return "the formula uses an alignment marker outside an approved math environment"
    if re.search(r"(?<!\\)[%#]", formula):
        return "the formula contains an unsupported TeX comment or parameter marker"
    if re.search(r"(?<!\\)[_^]\s*$", formula):
        return "the formula ends with an incomplete subscript or superscript"
    return None


def _literal_formula_markdown(value: str) -> str:
    escaped = html.escape(value.strip() or "(empty formula OCR block)", quote=False)
    indented = "\n".join(f"    {line}" for line in escaped.splitlines())
    return "*Formula could not be safely rendered; raw OCR follows.*\n\n" + indented


def _safe_formula_markdown(value: str) -> tuple[str, str | None]:
    formula = _strip_formula_delimiters(value)
    rejection = _formula_rejection_reason(formula)
    if rejection:
        return _literal_formula_markdown(value), rejection
    return f"$$\n{formula}\n$$", None


def _accepted_display_math_count(markdown: str) -> int:
    return sum(
        _formula_rejection_reason(match.group(1).strip()) is None
        for match in _DISPLAY_MATH_BLOCK.finditer(markdown)
    )


def _clean_table_html(value: str) -> str:
    cleaned = _RAW_SCRIPT_TAG.sub("", value)
    cleaned = _RAW_IMAGE_TAG.sub("", cleaned)
    cleaned = _EVENT_ATTRIBUTE.sub("", cleaned)
    return cleaned.strip()


def _element_key(element: Mapping[str, Any]) -> str:
    return f"{int(element.get('page_index', 0))}:{int(element.get('index', 0))}"


def _element_markdown(
    element: Mapping[str, Any],
    *,
    corrected_content: str,
    figure_assets: Mapping[str, str],
    figure_number: int,
) -> tuple[str, int, str | None]:
    label = str(element.get("type", "text")).lower().replace("-", "_")
    page_number = int(element.get("page_index", 0)) + 1
    if any(kind in label for kind in ("image", "figure", "chart", "illustration", "graphic")):
        asset = figure_assets.get(_element_key(element))
        if not asset:
            return "", figure_number, None
        figure_number += 1
        context = re.sub(r"\s+", " ", str(element.get("caption_context", ""))).strip()
        visible_caption = context or f"Content figure from source page {page_number}"
        alt = _safe_figure_alt(
            visible_caption,
            fallback=f"Content figure from source page {page_number}",
        )
        labels = [str(value).strip() for value in element.get("recognized_labels", []) if str(value).strip()]
        label_text = ""
        if labels:
            label_text = "\n\n**Recognized figure labels:** " + "; ".join(
                _escape_plain(value) for value in labels
            )
        return (
            f"![{alt}]({asset})\n\n"
            f"*Figure {figure_number}. {_escape_plain(visible_caption)}.*{label_text}",
            figure_number,
            None,
        )
    if "formula" in label or "equation" in label:
        markdown, rejection = _safe_formula_markdown(corrected_content)
        return markdown, figure_number, rejection
    if not corrected_content.strip():
        return "", figure_number, None
    if "table" in label:
        if "<table" in corrected_content.lower():
            return _clean_table_html(corrected_content), figure_number, None
        return _escape_plain(corrected_content), figure_number, None
    escaped = _escape_plain(corrected_content.strip())
    marks = [str(value) for value in element.get("margin_marks", []) if str(value)]
    mark_suffix = " " + " ".join(f"**[{value} marks]**" for value in marks) if marks else ""
    if label in {"doc_title", "document_title", "title"}:
        return f"### {escaped}", figure_number, None
    if "heading" in label or label in {"section_title", "sub_title", "subtitle"}:
        return f"### {escaped}{mark_suffix}", figure_number, None
    if label == "question":
        match = _QUESTION_START.match(corrected_content.strip())
        if match:
            prefix = _escape_plain(match.group(0).strip())
            body = _escape_plain(corrected_content.strip()[match.end() :].lstrip())
            return f"**{prefix}** {body}{mark_suffix}".strip(), figure_number, None
        return f"{escaped}{mark_suffix}", figure_number, None
    if "list" in label:
        items = []
        for line in corrected_content.splitlines():
            line = line.strip()
            if not line:
                continue
            if re.match(r"^(?:[-*+]\s+|\d+[.)]\s+)", line):
                items.append(_escape_plain(line))
            else:
                items.append(f"- {_escape_plain(line)}")
        return "\n".join(items), figure_number, None
    return f"{escaped}{mark_suffix}", figure_number, None


def _text_elements(result: Mapping[str, Any]) -> list[dict[str, Any]]:
    elements: list[dict[str, Any]] = []
    for index, line in enumerate(result.get("lines", [])):
        if not isinstance(line, Mapping):
            continue
        elements.append(
            {
                "type": "text",
                "bbox": line.get("bbox", []),
                "normalized_bbox": line.get("normalized_bbox", []),
                "content": str(line.get("retry_text", line.get("text", ""))),
                "confidence": float(
                    line.get("retry_confidence", line.get("confidence", 0.0)) or 0.0
                ),
                "index": index,
                "page_index": int(line.get("page_index", 0) or 0),
                "line_id": str(
                    line.get(
                        "line_id",
                        f"p{int(line.get('page_index', 0) or 0) + 1:04d}-l{index + 1:05d}",
                    )
                ),
            }
        )
    return elements


def _bbox_overlap_ratio(inner: Any, outer: Any) -> float:
    return geometry_overlap_ratio(inner, outer)


def _bbox_center_y(raw: Any) -> float:
    bbox = _bbox_values(raw)
    return (bbox[1] + bbox[3]) / 2 if bbox else float("inf")


def _merge_text_anchors(
    text_elements: Sequence[Mapping[str, Any]],
    anchors: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    merged = [dict(item) for item in text_elements]
    for anchor in sorted(anchors, key=_bbox_center_y):
        anchor_y = _bbox_center_y(anchor)
        position = len(merged)
        for index, item in enumerate(merged):
            if _bbox_center_y(item) > anchor_y:
                position = index
                break
        merged.insert(position, dict(anchor))
    return merged


def _markdown_cell(value: Any) -> str:
    return re.sub(r"\s*\n\s*", "<br>", str(value)).replace("|", r"\|").strip()


def _normalize_table_numeric_punctuation(value: str) -> tuple[str, bool]:
    """Remove isolated scan specks around integer-like table cells, not decimal values."""
    text = re.sub(r"\s+", " ", value).strip()
    updated = re.sub(r"^[.。·](?=\d{3,}$)", "", text)
    updated = re.sub(r"(?<=\d{3})[.。·]$", "", updated)
    updated = re.sub(r"(?<=\d)\s+[。·]$", "", updated)
    return updated, updated != value


def _cells_to_public_markdown(cells: Sequence[Mapping[str, Any]]) -> str:
    from app.workflows.table_extract import _cells_to_grid

    normalized = [dict(cell) for cell in cells]
    if not normalized:
        return ""
    has_spans = any(
        int(cell.get("row_span", 1)) > 1 or int(cell.get("col_span", 1)) > 1
        for cell in normalized
    )
    if not has_spans:
        grid = _cells_to_grid(normalized)
        if not grid:
            return ""
        header = "| " + " | ".join(_markdown_cell(value) for value in grid[0]) + " |"
        separator = "| " + " | ".join("---" for _ in grid[0]) + " |"
        body = [
            "| " + " | ".join(_markdown_cell(value) for value in row) + " |"
            for row in grid[1:]
        ]
        return "\n".join([header, separator, *body])

    row_count = max(int(cell.get("row", 0)) + int(cell.get("row_span", 1)) for cell in normalized)
    by_position = {
        (int(cell.get("row", 0)), int(cell.get("col", 0))): cell for cell in normalized
    }
    occupied: set[tuple[int, int]] = set()
    rows = ["<table>"]
    for row in range(row_count):
        rows.append("  <tr>")
        for (cell_row, cell_col), cell in sorted(by_position.items()):
            if cell_row != row or (cell_row, cell_col) in occupied:
                continue
            row_span = int(cell.get("row_span", 1))
            col_span = int(cell.get("col_span", 1))
            tag = "th" if row == 0 else "td"
            attributes = ""
            if row_span > 1:
                attributes += f' rowspan="{row_span}"'
            if col_span > 1:
                attributes += f' colspan="{col_span}"'
            value = html.escape(str(cell.get("text", "")).strip())
            rows.append(f"    <{tag}{attributes}>{value}</{tag}>")
            for occupied_row in range(cell_row, cell_row + row_span):
                for occupied_col in range(cell_col, cell_col + col_span):
                    occupied.add((occupied_row, occupied_col))
        rows.append("  </tr>")
    rows.append("</table>")
    return "\n".join(rows)


def _page_count(result: Mapping[str, Any], elements: Sequence[Mapping[str, Any]]) -> int:
    candidate = result.get("page_count")
    if isinstance(candidate, int) and candidate >= 1:
        return candidate
    pages = [int(item.get("page_index", 0) or 0) for item in elements]
    tables = result.get("tables", [])
    if isinstance(tables, Sequence):
        pages.extend(
            int(item.get("page_index", 0) or 0)
            for item in tables
            if isinstance(item, Mapping)
        )
    return max(pages, default=0) + 1


_QUESTION_START = re.compile(r"^(?:\d{1,3}\s*[.)]|\(?[a-z]\)|\([ivxlcdm]+\))\s*", re.I)
_EXAM_SIGNAL = re.compile(
    r"(?:subject\s*code|full\s*marks|examination|answer\s+any|\b\d+\s*[.)]\s*\(?[a-z]\)?)",
    re.I,
)


def _resolve_document_profile(elements: Sequence[Mapping[str, Any]], requested: str) -> str:
    selected = validate_document_profile(requested)
    if selected != "auto":
        return selected
    first_pages = "\n".join(
        str(item.get("content", ""))
        for item in elements
        if int(item.get("page_index", 0) or 0) <= 1
    )
    signals = _EXAM_SIGNAL.findall(first_pages)
    subquestions = len(re.findall(r"(?:^|\s)\([a-z]\)\s", first_pages, re.I))
    return "exam" if len(signals) + min(subquestions, 3) >= 3 else "general"


def _header_consensus(
    elements: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[Correction]]:
    candidates = [
        item
        for item in elements
        if item.get("type") == "text"
        and (_bbox_values(item) or (0, 1, 0, 1))[1] < 0.25
        and 8 <= len(str(item.get("content", "")).strip()) <= 100
    ]
    audits: list[Correction] = []
    for item in candidates:
        original = str(item.get("content", "")).strip()
        similar = [
            other
            for other in candidates
            if int(other.get("page_index", 0)) != int(item.get("page_index", 0))
            and SequenceMatcher(
                None,
                re.sub(r"\W+", "", original).casefold(),
                re.sub(r"\W+", "", str(other.get("content", ""))).casefold(),
            ).ratio()
            >= 0.84
        ]
        frequencies: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for other in similar:
            frequencies[str(other.get("content", "")).strip()].append(other)
        if not frequencies:
            continue
        canonical, matches = max(frequencies.items(), key=lambda entry: len(entry[1]))
        if canonical == original or len(matches) < 2:
            continue
        canonical_confidence = max(float(match.get("confidence", 0.0) or 0.0) for match in matches)
        original_confidence = float(item.get("confidence", 0.0) or 0.0)
        if canonical_confidence < original_confidence:
            continue
        item["content"] = canonical
        audits.append(
            Correction(
                original=original,
                corrected=canonical,
                page=int(item.get("page_index", 0)) + 1,
                block=int(item.get("index", 0)),
                kind="cross_page_consensus",
                confidence=round(min(0.99, canonical_confidence), 4),
                reason=["repeated high-confidence header text on other source pages"],
            )
        )
    return elements, audits


def _join_text_lines(first: Mapping[str, Any], second: Mapping[str, Any]) -> bool:
    if first.get("type") not in {"text", "question"} or second.get("type") != "text":
        return False
    next_text = str(second.get("content", "")).strip()
    if not next_text or _QUESTION_START.match(next_text):
        return False
    first_box, second_box = _bbox_values(first), _bbox_values(second)
    if first_box is None or second_box is None:
        return False
    first_height = max(0.001, first_box[3] - first_box[1])
    vertical_gap = second_box[1] - first_box[3]
    indentation = abs(second_box[0] - first_box[0])
    return -first_height <= vertical_gap <= first_height * 1.8 and indentation <= 0.14


def _prepare_text_page(
    text_elements: Sequence[Mapping[str, Any]],
    *,
    profile: str,
) -> list[dict[str, Any]]:
    items = [dict(item) for item in text_elements]
    for item in items:
        item["line_ids"] = [str(item.get("line_id"))] if item.get("line_id") else []
    if profile == "exam":
        marks = [
            item
            for item in items
            if re.fullmatch(r"\d{1,2}", str(item.get("content", "")).strip())
            and (_bbox_values(item) or (0, 0, 0, 0))[0] >= 0.80
            and (_bbox_values(item) or (0, 0, 0, 1))[3] <= 0.95
        ]
        for mark in marks:
            mark_box = _bbox_values(mark)
            if mark_box is None:
                continue
            mark_y = (mark_box[1] + mark_box[3]) / 2
            targets = [
                item
                for item in items
                if item is not mark
                and item not in marks
                and _bbox_values(item) is not None
                and abs(
                    ((_bbox_values(item) or (0, 0, 0, 0))[1]
                    + (_bbox_values(item) or (0, 0, 0, 0))[3])
                    / 2
                    - mark_y
                )
                <= 0.025
            ]
            if not targets:
                continue
            target = min(
                targets,
                key=lambda item: abs(
                    ((_bbox_values(item) or (0, 0, 0, 0))[1]
                    + (_bbox_values(item) or (0, 0, 0, 0))[3])
                    / 2
                    - mark_y
                ),
            )
            target.setdefault("margin_marks", []).append(str(mark.get("content", "")).zfill(2))
            target["line_ids"].extend(mark.get("line_ids", []))
            items.remove(mark)

        first_question_y = min(
            (
                (_bbox_values(item) or (0, 1, 0, 1))[1]
                for item in items
                if _QUESTION_START.match(str(item.get("content", "")).strip())
            ),
            default=0.28,
        )
        for item in items:
            box = _bbox_values(item)
            text = str(item.get("content", "")).strip()
            if _QUESTION_START.match(text):
                item["type"] = "question"
            elif box and box[1] < min(0.25, first_question_y) and len(text) <= 100:
                item["type"] = "heading"

    merged: list[dict[str, Any]] = []
    for item in items:
        if merged and _join_text_lines(merged[-1], item):
            prior = merged[-1]
            prior["content"] = f"{str(prior.get('content', '')).rstrip()} {str(item.get('content', '')).lstrip()}"
            prior["line_ids"].extend(item.get("line_ids", []))
            prior["confidence"] = min(
                float(prior.get("confidence", 0.0) or 0.0),
                float(item.get("confidence", 0.0) or 0.0),
            )
            prior["bbox"] = item.get("bbox", prior.get("bbox"))
            prior["normalized_bbox"] = item.get(
                "normalized_bbox", prior.get("normalized_bbox")
            )
        else:
            merged.append(item)
    return merged


def build_structured_markdown(
    result: Mapping[str, Any],
    *,
    workflow: str,
    source_name: str,
    output_stem: str,
    figure_assets: Mapping[str, str],
    settings: Settings | None = None,
    supplemental_elements: Sequence[Mapping[str, Any]] | None = None,
) -> tuple[
    str,
    list[dict[str, Any]],
    list[dict[str, Any]],
    list[str],
    list[dict[str, Any]],
]:
    """Build page-aligned Markdown while retaining the raw workflow fields separately."""
    cfg = settings or get_settings()
    raw_elements = result.get("elements")
    elements = (
        [dict(item) for item in raw_elements if isinstance(item, Mapping)]
        if isinstance(raw_elements, Sequence) and not isinstance(raw_elements, (str, bytes))
        else _text_elements(result)
    )
    supplemental = [dict(item) for item in supplemental_elements or ()]
    requested_profile = str(result.get("requested_document_profile", "auto") or "auto")
    selected_profile = _resolve_document_profile(elements, requested_profile)
    consensus_audits: list[Correction] = []
    if workflow == "text_recognition":
        elements, consensus_audits = _header_consensus(elements)
        supplemental = [
            item
            for item in supplemental
            if any(
                kind in str(item.get("type", "")).lower()
                for kind in _TEXT_SUPPLEMENTAL_TYPES
            )
        ]

    raw_tables = result.get("tables")
    table_elements: list[dict[str, Any]] = []
    table_pages: set[int] = set()
    if isinstance(raw_tables, Sequence) and not isinstance(raw_tables, (str, bytes)):
        for position, table in enumerate(raw_tables):
            if not isinstance(table, Mapping):
                continue
            page_index = int(table.get("page_index", 0) or 0)
            table_pages.add(page_index)
            table_elements.append(
                {
                    "type": "structured_table",
                    "bbox": table.get("bbox", []),
                    "normalized_bbox": table.get("normalized_bbox", []),
                    "content": "",
                    "index": 2_000_000 + position,
                    "page_index": page_index,
                    "table": table,
                }
            )

    if workflow == "text_recognition":
        dispositions = {
            str(item.get("line_id")): (
                "prose" if str(item.get("content", "")).strip() else "unresolved"
            )
            for item in elements
            if item.get("line_id")
        }
        accepted_table_lines: set[str] = set()
        for table_element in table_elements:
            table = table_element.get("table")
            if not isinstance(table, Mapping) or table.get("status") != "accepted":
                continue
            accepted_table_lines.update(
                str(value) for value in table.get("source_line_ids", ()) if value
            )
        for line_id in accepted_table_lines:
            if line_id in dispositions:
                dispositions[line_id] = "table_cell"

        deduplicated_supplemental: list[dict[str, Any]] = []
        formulas: list[dict[str, Any]] = []
        for item in supplemental:
            label = str(item.get("type", "")).lower()
            if not any(kind in label for kind in ("formula", "equation")):
                deduplicated_supplemental.append(item)
                continue
            content_key = re.sub(r"\s+", "", str(item.get("content", ""))).casefold()
            duplicate = any(
                content_key
                and int(item.get("page_index", 0) or 0)
                == int(other.get("page_index", 0) or 0)
                and content_key
                == re.sub(r"\s+", "", str(other.get("content", ""))).casefold()
                and _bbox_iou(item, other) >= 0.65
                for other in formulas
            )
            if not duplicate:
                formulas.append(item)
                deduplicated_supplemental.append(item)
        supplemental = deduplicated_supplemental

        for formula in formulas:
            content = str(formula.get("content", ""))
            if _formula_rejection_reason(_strip_formula_delimiters(content)) is not None:
                continue
            source_lines: list[dict[str, Any]] = []
            for item in elements:
                line_id = str(item.get("line_id", ""))
                if (
                    not line_id
                    or int(item.get("page_index", 0) or 0)
                    != int(formula.get("page_index", 0) or 0)
                    or line_id in accepted_table_lines
                    or _bbox_overlap_ratio(item, formula) < 0.55
                ):
                    continue
                dispositions[line_id] = "verified_duplicate"
                source_lines.append(
                    {
                        "line_id": line_id,
                        "text": str(item.get("content", "")),
                        "bbox": item.get("normalized_bbox") or item.get("bbox", []),
                    }
                )
            formula["source_line_ids"] = [line["line_id"] for line in source_lines]
            formula["raw_ocr_lines"] = source_lines

        accepted_figures = [
            item
            for item in supplemental
            if item.get("figure_status") == "accepted" and _element_key(item) in figure_assets
        ]
        for figure in accepted_figures:
            labels: list[str] = []
            label_ids: list[str] = []
            for item in elements:
                line_id = str(item.get("line_id", ""))
                if (
                    not line_id
                    or int(item.get("page_index", 0) or 0)
                    != int(figure.get("page_index", 0) or 0)
                    or line_id in accepted_table_lines
                ):
                    continue
                if _bbox_overlap_ratio(item, figure) >= 0.70:
                    labels.append(str(item.get("content", "")).strip())
                    label_ids.append(line_id)
                    dispositions[line_id] = "figure_label"
            figure["recognized_labels"] = list(dict.fromkeys(value for value in labels if value))
            figure["source_line_ids"] = label_ids

        visible_text = [
            item
            for item in elements
            if dispositions.get(str(item.get("line_id", "")), "prose") == "prose"
        ]
        anchors = [*supplemental, *table_elements]
        by_page_text: dict[int, list[dict[str, Any]]] = defaultdict(list)
        by_page_anchor: dict[int, list[dict[str, Any]]] = defaultdict(list)
        for item in visible_text:
            by_page_text[int(item.get("page_index", 0) or 0)].append(item)
        for item in anchors:
            by_page_anchor[int(item.get("page_index", 0) or 0)].append(item)
        elements = []
        for page_index in sorted(set(by_page_text) | set(by_page_anchor)):
            prepared = _prepare_text_page(
                by_page_text[page_index],
                profile=selected_profile,
            )
            for anchor in by_page_anchor[page_index]:
                if not any(
                    term in str(anchor.get("type", "")).lower()
                    for term in ("figure", "image", "chart")
                ):
                    continue
                anchor_y = _bbox_center_y(anchor)
                preceding = [item for item in prepared if _bbox_center_y(item) < anchor_y]
                if preceding:
                    anchor["caption_context"] = str(preceding[-1].get("content", ""))[:180]
            elements.extend(_merge_text_anchors(prepared, by_page_anchor[page_index]))

        raw_line_count = len(dispositions)
        accounted = sum(
            value in {"prose", "table_cell", "figure_label", "verified_duplicate"}
            for value in dispositions.values()
        )
        coverage = {
            "status": "passed" if accounted == raw_line_count else "failed",
            "raw_line_count": raw_line_count,
            "emitted_prose_count": sum(value == "prose" for value in dispositions.values()),
            "verified_table_line_count": sum(
                value == "table_cell" for value in dispositions.values()
            ),
            "figure_label_line_count": sum(
                value == "figure_label" for value in dispositions.values()
            ),
            "verified_duplicate_count": sum(
                value == "verified_duplicate" for value in dispositions.values()
            ),
            "user_removed_count": 0,
            "unaccounted_line_count": raw_line_count - accounted,
            "coverage_percent": round(100.0 * accounted / raw_line_count, 2)
            if raw_line_count
            else 100.0,
            "line_dispositions": dispositions,
        }
        if isinstance(result, dict):
            result["document_profile"] = selected_profile
            result["content_coverage"] = coverage
            result["structured_formulas"] = [dict(formula) for formula in formulas]
            raw_lines = result.get("lines")
            if isinstance(raw_lines, list):
                for line in raw_lines:
                    if isinstance(line, dict) and line.get("line_id") in dispositions:
                        line["disposition"] = dispositions[str(line["line_id"])]
    else:
        known = {_element_key(item) for item in elements}
        elements.extend(item for item in supplemental if _element_key(item) not in known)
        if table_pages:
            elements = [
                item
                for item in elements
                if not (
                    "table" in str(item.get("type", "")).lower()
                    and int(item.get("page_index", 0) or 0) in table_pages
                )
            ]
        elements.extend(table_elements)
        elements = suppress_adjacent_duplicates(elements)
        if isinstance(result, dict):
            result["document_profile"] = selected_profile
    pages: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for element in elements:
        pages[int(element.get("page_index", 0) or 0)].append(element)

    corrector = ConservativeCorrector(cfg.correction_enabled)
    corrections: list[dict[str, Any]] = [audit.as_dict() for audit in consensus_audits]
    warnings: list[str] = []
    table_specs: list[dict[str, Any]] = []
    parts = [f"# {_escape_plain(output_stem)}", "", f"Source file: `{_escape_plain(source_name)}`"]
    figure_number = 0
    count = _page_count(result, elements)

    def corrected_table_cells(
        cells: Sequence[Mapping[str, Any]],
        *,
        page_number: int,
        table_block: int,
    ) -> list[dict[str, Any]]:
        normalized: list[dict[str, Any]] = []
        for offset, cell in enumerate(cells):
            copied = dict(cell)
            content = str(copied.get("text", ""))
            corrected, audit = corrector.correct(
                content,
                page=page_number,
                block=table_block + offset,
            )
            corrected, numeric_punctuation_changed = _normalize_table_numeric_punctuation(
                corrected
            )
            copied["text"] = corrected
            if audit:
                corrections.append(audit.as_dict())
            if numeric_punctuation_changed:
                corrections.append(
                    Correction(
                        original=content,
                        corrected=corrected,
                        page=page_number,
                        block=table_block + offset,
                        kind="table_numeric_punctuation",
                        confidence=0.99,
                        reason=["isolated OCR punctuation removed from a numeric table cell"],
                    ).as_dict()
                )
            normalized.append(copied)
        rows: dict[int, list[dict[str, Any]]] = defaultdict(list)
        for cell in normalized:
            rows[int(cell.get("row", 0) or 0)].append(cell)
        for row_cells in rows.values():
            integer_like = [
                cell
                for cell in row_cells
                if re.fullmatch(r"[.·。]?\d{2,}", str(cell.get("text", "")).strip())
            ]
            if len(integer_like) < 3:
                continue
            for cell in integer_like:
                value = str(cell.get("text", "")).strip()
                match = re.fullmatch(r"[.·。](\d{2,})", value)
                if not match:
                    continue
                cell["text"] = match.group(1)
                corrections.append(
                    Correction(
                        original=value,
                        corrected=match.group(1),
                        page=page_number,
                        block=table_block + int(cell.get("col", 0) or 0),
                        kind="table_numeric_punctuation",
                        confidence=0.96,
                        reason=["leading scan speck removed from an otherwise integer-valued row"],
                    ).as_dict()
                )
        return normalized

    for page_index in range(count):
        parts.extend(["", f"## Source page {page_index + 1}", ""])
        page_had_content = False
        page_elements = pages.get(page_index, [])
        if workflow != "text_recognition":
            page_elements = sorted(page_elements, key=_reading_key)
        for element in page_elements:
            label = str(element.get("type", "")).lower()
            if label == "structured_table":
                table = element.get("table")
                cells = table.get("cells") if isinstance(table, Mapping) else None
                table_status = str(table.get("status", "unverified")) if isinstance(table, Mapping) else "unverified"
                if isinstance(cells, Sequence) and not isinstance(cells, (str, bytes)) and cells:
                    valid_cells = [cell for cell in cells if isinstance(cell, Mapping)]
                    valid_cells = corrected_table_cells(
                        valid_cells,
                        page_number=page_index + 1,
                        table_block=int(element.get("index", 0) or 0) * 1_000,
                    )
                    if isinstance(table, dict):
                        table["cells"] = valid_cells
                    public_markdown = _cells_to_public_markdown(
                        valid_cells
                    )
                    if public_markdown:
                        table_index = int(table.get("table_index", 0)) + 1
                        parts.extend([f"### Table {table_index}", ""])
                        if table_status != "accepted":
                            parts.extend(
                                [
                                    "> **Needs review:** this table could not be fully verified against primary OCR.",
                                    "",
                                ]
                            )
                            review_image = str(table.get("review_image", "")).strip()
                            if review_image:
                                parts.extend(
                                    [
                                        f"![Unverified table {table_index} from source page "
                                        f"{page_index + 1}]({review_image})",
                                        "",
                                    ]
                                )
                            evidence = list(
                                dict.fromkeys(
                                    str(cell.get("source_text", "")).strip()
                                    for cell in valid_cells
                                    if str(cell.get("source_text", "")).strip()
                                )
                            )
                            if evidence:
                                parts.extend(
                                    [
                                        "**Primary OCR evidence:**",
                                        "",
                                        *[f"- {_escape_plain(value)}" for value in evidence],
                                        "",
                                    ]
                                )
                        else:
                            marker = f"[[OCR_TABLE_{len(table_specs) + 1:03d}]]"
                            table_specs.append(
                                {
                                    "marker": marker,
                                    "cells": valid_cells,
                                    "public_markdown": public_markdown,
                                    "status": table_status,
                                }
                            )
                            parts.extend([public_markdown, ""])
                        page_had_content = True
                continue
            if workflow == "table_extraction" and "table" in label:
                continue
            content = str(element.get("content", ""))
            block = int(element.get("index", 0) or 0)
            confidence = element.get("confidence")
            if (
                workflow == "text_recognition"
                and isinstance(confidence, (int, float))
                and float(confidence) < cfg.low_confidence_threshold
                and content.strip()
            ):
                corrector._add_review(
                    page=page_index + 1,
                    block=block,
                    original=content,
                    suggested=None,
                    kind="low_confidence",
                    confidence=float(confidence),
                    reason="OCR confidence is below the automatic-review threshold",
                )
            if any(term in label for term in ("formula", "equation", "code", "table")):
                corrected = content
                audit = None
            else:
                corrected, audit = corrector.correct(
                    content,
                    page=page_index + 1,
                    block=block,
                )
            if audit:
                corrections.append(audit.as_dict())
            markdown, figure_number, formula_rejection = _element_markdown(
                element,
                corrected_content=corrected,
                figure_assets=figure_assets,
                figure_number=figure_number,
            )
            if formula_rejection:
                corrector._add_review(
                    page=page_index + 1,
                    block=block,
                    original=content,
                    suggested=None,
                    kind="formula_rendering",
                    confidence=0.0,
                    reason=formula_rejection,
                )
                warnings.append(
                    f"Formula on source page {page_index + 1}, block {block} was kept "
                    f"as literal text because {formula_rejection}."
                )
            if "table" in label and "<table" in content.lower():
                try:
                    from app.workflows.table_extract import _parse_html_cells

                    cells = _parse_html_cells(content, cfg.max_table_cells)
                    cells = corrected_table_cells(
                        cells,
                        page_number=page_index + 1,
                        table_block=block * 1_000,
                    )
                    public_markdown = _cells_to_public_markdown(cells)
                    marker = f"[[OCR_TABLE_{len(table_specs) + 1:03d}]]"
                    table_specs.append(
                        {
                            "marker": marker,
                            "cells": cells,
                            "public_markdown": public_markdown,
                        }
                    )
                    markdown = public_markdown
                except Exception:
                    warnings.append(
                        f"A table on source page {page_index + 1} could not be structured; "
                        "its visible text was retained."
                    )
                    visible = re.sub(r"<[^>]+>", " ", content)
                    markdown = _escape_plain(re.sub(r"\s+", " ", visible).strip())
            if markdown:
                parts.extend([markdown, ""])
                page_had_content = True

        if not page_had_content:
            parts.extend(["*No readable content was detected on this source page.*", ""])

    if corrector.warning:
        warnings.append(corrector.warning)
    markdown = "\n".join(parts).rstrip() + "\n"
    return markdown, corrections, corrector.review_suggestions, warnings, table_specs


def _set_run_font(run: Any, name: str, size: float, color: str = "000000") -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def _set_style(style: Any, *, size: float, color: str, before: float, after: float, line: float) -> None:
    from docx.shared import Pt, RGBColor

    style.font.name = "Carlito"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def create_reference_docx(path: Path, source_name: str) -> None:
    """Create the compact_reference_guide preset with the requested A4 override."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Mm, Pt

    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = document.styles
    _set_style(styles["Normal"], size=11, color="000000", before=0, after=6, line=1.25)
    _set_style(styles["Title"], size=22, color="0B2545", before=0, after=12, line=1.0)
    _set_style(styles["Subtitle"], size=11, color="64748B", before=0, after=8, line=1.0)
    _set_style(styles["Heading 1"], size=16, color="2E74B5", before=18, after=10, line=1.0)
    _set_style(styles["Heading 2"], size=13, color="2E74B5", before=14, after=7, line=1.0)
    _set_style(styles["Heading 3"], size=12, color="1F4D78", before=10, after=5, line=1.0)
    _set_style(styles["Caption"], size=9, color="64748B", before=4, after=8, line=1.0)
    styles["Caption"].font.italic = True
    for name in ("List Bullet", "List Number"):
        _set_style(styles[name], size=11, color="000000", before=0, after=4, line=1.25)
        styles[name].paragraph_format.left_indent = Inches(0.375)
        styles[name].paragraph_format.first_line_indent = Inches(-0.188)

    header = section.header.paragraphs[0]
    header.text = safe_output_stem(source_name, max_length=90)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        _set_run_font(run, "Carlito", 8.5, "64748B")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ")
    _append_page_field(footer)
    for run in footer.runs:
        _set_run_font(run, "Carlito", 8.5, "64748B")

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _append_page_field(paragraph: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)


def _set_cell_margins(cell: Any, *, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table: Any, total_width_dxa: int) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    column_count = len(table.columns)
    if not column_count:
        return
    weights: list[int] = []
    for col_index in range(column_count):
        lengths = [
            max((len(line) for line in row.cells[col_index].text.splitlines()), default=1)
            for row in table.rows
            if col_index < len(row.cells)
        ]
        weights.append(max(5, min(60, max(lengths, default=5))))
    weight_total = sum(weights)
    widths = [max(720, round(total_width_dxa * weight / weight_total)) for weight in weights]
    widths[-1] += total_width_dxa - sum(widths)

    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    table_width = tbl_pr.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        tbl_pr.append(table_width)
    table_width.set(qn("w:w"), str(total_width_dxa))
    table_width.set(qn("w:type"), "dxa")
    indent = tbl_pr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)

    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            width = widths[min(col_index, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_width = tc_pr.first_child_found_in("w:tcW")
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                tc_pr.append(tc_width)
            tc_width.set(qn("w:w"), str(width))
            tc_width.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = 0
                paragraph.paragraph_format.space_after = 0
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    _set_run_font(run, "Carlito", 9.5, "000000")
                    if row_index == 0:
                        run.bold = True
        if row_index == 0:
            row_pr = row._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            row_pr.append(repeat)
            for cell in row.cells:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8EEF5")
                cell._tc.get_or_add_tcPr().append(shading)


def _insert_structured_tables(document: Any, table_specs: Sequence[Mapping[str, Any]]) -> None:
    for spec in table_specs:
        marker = str(spec.get("marker", ""))
        cells = spec.get("cells")
        if not marker or not isinstance(cells, Sequence) or isinstance(cells, (str, bytes)):
            continue
        valid_cells = [item for item in cells if isinstance(item, Mapping)]
        if not valid_cells:
            continue
        row_count = max(int(item.get("row", 0)) + int(item.get("row_span", 1)) for item in valid_cells)
        column_count = max(
            int(item.get("col", 0)) + int(item.get("col_span", 1)) for item in valid_cells
        )
        placeholder = next(
            (paragraph for paragraph in document.paragraphs if paragraph.text.strip() == marker),
            None,
        )
        if placeholder is None:
            continue
        table = document.add_table(rows=row_count, cols=column_count)
        table.style = "Table Grid"
        for item in valid_cells:
            row = int(item.get("row", 0))
            column = int(item.get("col", 0))
            row_span = int(item.get("row_span", 1))
            column_span = int(item.get("col_span", 1))
            start = table.cell(row, column)
            end = table.cell(row + row_span - 1, column + column_span - 1)
            target = start.merge(end) if row_span > 1 or column_span > 1 else start
            target.text = str(item.get("text", ""))
        placeholder._p.addnext(table._tbl)
        placeholder._element.getparent().remove(placeholder._element)


def _postprocess_docx(
    path: Path,
    source_name: str,
    table_specs: Sequence[Mapping[str, Any]] = (),
) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt

    document = Document(path)
    for section in document.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)
        header = section.header.paragraphs[0]
        header.text = safe_output_stem(source_name, max_length=90)
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in header.runs:
            _set_run_font(run, "Carlito", 8.5, "64748B")
        footer = section.footer.paragraphs[0]
        footer.clear()
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Page ")
        _append_page_field(footer)
        for run in footer.runs:
            _set_run_font(run, "Carlito", 8.5, "64748B")

    _insert_structured_tables(document, table_specs)

    for paragraph in document.paragraphs:
        marker = _PAGE_MARKER.match(paragraph.text.strip())
        if marker:
            paragraph.paragraph_format.keep_with_next = True
            if int(marker.group(1)) > 1:
                page_break = OxmlElement("w:pageBreakBefore")
                page_break.set(qn("w:val"), "true")
                paragraph._p.get_or_add_pPr().append(page_break)
        if paragraph.text.strip().startswith("Figure ") and "Extracted from source page" in paragraph.text:
            paragraph.style = document.styles["Caption"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if paragraph.text.strip().startswith("Recognized figure labels:"):
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 0.95
            for run in paragraph.runs:
                _set_run_font(run, "Carlito", 9)

    content_width_dxa = round(
        (
            document.sections[0].page_width
            - document.sections[0].left_margin
            - document.sections[0].right_margin
        )
        / 635
    )
    for table in document.tables:
        _set_table_geometry(table, content_width_dxa)

    for index, shape in enumerate(document.inline_shapes, start=1):
        max_width = (
            document.sections[0].page_width
            - document.sections[0].left_margin
            - document.sections[0].right_margin
        )
        if shape.width > max_width:
            ratio = max_width / shape.width
            shape.width = max_width
            shape.height = round(shape.height * ratio)
        max_figure_height = Inches(2.25)
        if shape.height > max_figure_height:
            ratio = max_figure_height / shape.height
            shape.height = max_figure_height
            shape.width = round(shape.width * ratio)
        doc_pr = shape._inline.docPr
        if not doc_pr.get("descr"):
            doc_pr.set("descr", f"Extracted source figure {index}")
        paragraph = shape._inline.getparent().getparent()
        p_pr = paragraph.find(qn("w:pPr"))
        if p_pr is None:
            p_pr = OxmlElement("w:pPr")
            paragraph.insert(0, p_pr)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), str(round(Pt(4) / 635)))
        spacing.set(qn("w:after"), str(round(Pt(4) / 635)))
        p_pr.append(spacing)

    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        document.save(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _validate_local_resources(markdown: str, bundle_root: Path) -> None:
    root = bundle_root.resolve()
    for match in _IMAGE_MARKDOWN.finditer(markdown):
        _validate_local_resource_target(match.group(1), root)


def _validate_local_resource_target(target_value: str, root: Path) -> None:
    target = target_value.strip().split(maxsplit=1)[0].strip("<>")
    lowered = target.casefold()
    if (
        "://" in target
        or target.startswith(("//", "\\\\"))
        or lowered.startswith(("data:", "file:"))
    ):
        raise ValueError("Remote or embedded image resources are not permitted")
    candidate = (root / target).resolve()
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise ValueError("Image resource escapes the document bundle") from exc
    if not candidate.is_file():
        raise ValueError(f"Generated image resource is missing: {target}")


def _validate_pandoc_ast(
    markdown_path: Path,
    *,
    bundle_root: Path,
    timeout_seconds: int,
) -> None:
    """Reject links and nonlocal image nodes before DOCX conversion."""
    command = [
        str(pandoc_executable()),
        str(markdown_path),
        "--from=markdown+raw_html+raw_attribute+tex_math_dollars"
        "-implicit_figures-autolink_bare_uris",
        "--to=json",
        f"--resource-path={bundle_root}",
    ]
    try:
        completed = subprocess.run(  # noqa: S603
            command,
            check=False,
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=timeout_seconds,
            cwd=bundle_root,
        )
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError("Markdown safety validation timed out") from exc
    if completed.returncode != 0:
        raise RuntimeError("Generated Markdown could not be safely parsed")
    try:
        ast = json.loads(completed.stdout)
    except (TypeError, json.JSONDecodeError) as exc:
        raise RuntimeError("Pandoc returned invalid document structure") from exc

    def visit(node: Any) -> None:
        if isinstance(node, dict):
            node_type = node.get("t")
            content = node.get("c")
            if node_type == "Link":
                raise RuntimeError("Generated Markdown contains an unintended hyperlink")
            if node_type == "Image":
                try:
                    target = str(content[2][0])
                except (IndexError, TypeError) as exc:
                    raise RuntimeError("Generated Markdown contains an invalid image") from exc
                _validate_local_resource_target(target, bundle_root.resolve())
            for value in node.values():
                visit(value)
        elif isinstance(node, list):
            for value in node:
                visit(value)

    visit(ast)


_RELATIONSHIP_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_OFFICE_RELATIONSHIP_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
_WORDPROCESSING_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _source_part_for_relationships(rels_name: str) -> str | None:
    if rels_name == "_rels/.rels":
        return None
    directory, filename = posixpath.split(rels_name)
    if posixpath.basename(directory) != "_rels" or not filename.endswith(".rels"):
        return None
    return posixpath.join(posixpath.dirname(directory), filename[:-5])


def _unwrap_external_hyperlinks(xml: bytes, relationship_ids: set[str]) -> bytes:
    if not relationship_ids:
        return xml
    from lxml import etree

    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    root = etree.fromstring(xml, parser=parser)
    relationship_key = f"{{{_OFFICE_RELATIONSHIP_NS}}}id"
    hyperlink_tag = f"{{{_WORDPROCESSING_NS}}}hyperlink"
    for hyperlink in list(root.iter(hyperlink_tag)):
        if hyperlink.get(relationship_key) not in relationship_ids:
            continue
        parent = hyperlink.getparent()
        if parent is None:
            continue
        index = parent.index(hyperlink)
        for child in list(hyperlink):
            hyperlink.remove(child)
            parent.insert(index, child)
            index += 1
        parent.remove(hyperlink)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _sanitize_external_hyperlinks(path: Path) -> int:
    """Remove external hyperlink relationships while retaining their visible runs."""
    from lxml import etree

    with zipfile.ZipFile(path) as archive:
        entries = {info.filename: (info, archive.read(info.filename)) for info in archive.infolist()}

    removed = 0
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    for rels_name, (info, payload) in list(entries.items()):
        if not rels_name.endswith(".rels"):
            continue
        root = etree.fromstring(payload, parser=parser)
        hyperlink_ids: set[str] = set()
        changed = False
        for relationship in list(root):
            if str(relationship.get("TargetMode", "")).casefold() != "external":
                continue
            relation_type = str(relationship.get("Type", ""))
            if not relation_type.endswith("/hyperlink"):
                raise RuntimeError("The generated DOCX contains a forbidden external resource")
            hyperlink_ids.add(str(relationship.get("Id", "")))
            root.remove(relationship)
            removed += 1
            changed = True
        if not changed:
            continue
        entries[rels_name] = (
            info,
            etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
        )
        source_name = _source_part_for_relationships(rels_name)
        if source_name and source_name in entries:
            source_info, source_xml = entries[source_name]
            entries[source_name] = (
                source_info,
                _unwrap_external_hyperlinks(source_xml, hyperlink_ids),
            )

    if not removed:
        return 0
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.sanitize.tmp")
    try:
        with zipfile.ZipFile(temporary, "w") as archive:
            for info, payload in entries.values():
                archive.writestr(info, payload)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    return removed


def _docx_structure(path: Path) -> tuple[int, bool]:
    with zipfile.ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
        math_count = document_xml.count(b"<m:oMath")
        for name in archive.namelist():
            if not name.endswith(".rels"):
                continue
            relationships = archive.read(name)
            if b'TargetMode="External"' in relationships or b"TargetMode='External'" in relationships:
                return math_count, True
            if b"/image" in relationships:
                from lxml import etree

                root = etree.fromstring(relationships)
                source_part = _source_part_for_relationships(name)
                source_dir = posixpath.dirname(source_part or "")
                for relationship in root:
                    if not str(relationship.get("Type", "")).endswith("/image"):
                        continue
                    target = str(relationship.get("Target", ""))
                    resolved = posixpath.normpath(posixpath.join(source_dir, target))
                    if not resolved.startswith("word/media/"):
                        return math_count, True
    return math_count, False


def convert_markdown_to_docx(
    markdown_path: Path,
    docx_path: Path,
    *,
    source_name: str,
    settings: Settings | None = None,
    table_specs: Sequence[Mapping[str, Any]] = (),
) -> list[str]:
    """Convert local generated Markdown with pinned Pandoc, then normalize OOXML."""
    cfg = settings or get_settings()
    markdown = markdown_path.read_text(encoding="utf-8")
    _validate_local_resources(markdown, markdown_path.parent)
    reference = markdown_path.parent / ".compact-reference.docx"
    pandoc_markdown = markdown_path.with_name(
        f".{markdown_path.stem}.{uuid.uuid4().hex}.pandoc.md"
    )
    temporary = docx_path.with_name(f".{docx_path.name}.{uuid.uuid4().hex}.tmp.docx")
    warnings: list[str] = []
    internal_markdown = markdown
    for spec in table_specs:
        marker = str(spec.get("marker", ""))
        public_markdown = str(spec.get("public_markdown", ""))
        if not marker or not public_markdown or public_markdown not in internal_markdown:
            raise RuntimeError("A structured table could not be prepared for DOCX conversion")
        internal_markdown = internal_markdown.replace(public_markdown, marker, 1)
    try:
        pandoc_markdown.write_text(internal_markdown, encoding="utf-8")
        _validate_pandoc_ast(
            pandoc_markdown,
            bundle_root=markdown_path.parent,
            timeout_seconds=cfg.docx_timeout_seconds,
        )
        create_reference_docx(reference, source_name)
        command = [
            str(pandoc_executable()),
            str(pandoc_markdown),
            "--from=markdown+raw_html+raw_attribute+tex_math_dollars"
            "-implicit_figures-autolink_bare_uris",
            "--to=docx",
            f"--reference-doc={reference}",
            f"--resource-path={markdown_path.parent}",
            "--wrap=none",
            "--standalone",
            "--output",
            str(temporary),
        ]
        completed = subprocess.run(  # noqa: S603
            command,
            check=False,
            capture_output=True,
            text=True,
            timeout=cfg.docx_timeout_seconds,
            cwd=markdown_path.parent,
        )
        if completed.returncode != 0 or not temporary.is_file():
            message = (completed.stderr or completed.stdout or "Pandoc conversion failed").strip()
            raise RuntimeError(message[-500:])
        os.replace(temporary, docx_path)
        _postprocess_docx(docx_path, source_name, table_specs)
        removed_hyperlinks = _sanitize_external_hyperlinks(docx_path)
        if removed_hyperlinks:
            warnings.append(
                "One or more unsafe hyperlink relationships were removed while preserving "
                "their visible text."
            )
        math_expected = _accepted_display_math_count(markdown)
        math_count, has_external_relationship = _docx_structure(docx_path)
        if has_external_relationship:
            raise RuntimeError("The generated DOCX contains an external relationship")
        if math_count < math_expected:
            warnings.append(
                "One or more formulas could not be converted to native OMML equations; "
                "their LaTeX remains visible."
            )
        return warnings
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(
            f"DOCX generation exceeded the {cfg.docx_timeout_seconds}-second limit"
        ) from exc
    finally:
        reference.unlink(missing_ok=True)
        pandoc_markdown.unlink(missing_ok=True)
        temporary.unlink(missing_ok=True)

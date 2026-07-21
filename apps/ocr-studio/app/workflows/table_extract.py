"""PP-StructureV3 table extraction with bounded, injection-safe exports."""

from __future__ import annotations

import csv
import io
import logging
import math
import os
import re
import uuid
from collections import Counter
from collections.abc import Mapping, Sequence
from difflib import SequenceMatcher
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from PIL import Image

from app.core.capabilities import validate_document_profile, validate_language
from app.core.config import get_settings
from app.core.engine import get_manager
from app.core.errors import InputLimitError, WorkflowResultError
from app.core.geometry import (
    bbox_iou,
    bbox_values,
    normalized_bbox,
    overlap_ratio,
    union_bbox,
)
from app.core.progress import ProgressCallback, emit_progress
from app.core.serialization import atomic_write_json, atomic_write_text, json_safe, public_path
from app.workflows.pdf_utils import (
    clear_gpu_cache,
    is_pdf,
    iter_pdf_pages,
    translate_inference_error,
)
from app.workflows.result_adapters import parsing_elements, resolve_output_directory, result_payload

logger = logging.getLogger("predixalearn.workflow.table_extract")
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_MAX_TABLE_SKEW = 0.30
_SKEW_STEPS = 121


class _TableHTMLParser(HTMLParser):
    def __init__(self, max_cells: int) -> None:
        super().__init__(convert_charrefs=True)
        self.max_cells = max_cells
        self._cells: list[dict[str, Any]] = []
        self._occupied: set[tuple[int, int]] = set()
        self._row = -1
        self._col = 0
        self._in_cell = False
        self._text: list[str] = []
        self._row_span = 1
        self._col_span = 1

    @staticmethod
    def _span(attrs: Mapping[str, str | None], name: str) -> int:
        raw = attrs.get(name, "1") or "1"
        try:
            value = int(raw)
        except ValueError as exc:
            raise WorkflowResultError(f"Table HTML contains invalid {name}") from exc
        if value < 1:
            raise WorkflowResultError(f"Table HTML contains invalid {name}")
        return value

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag == "tr":
            self._row += 1
            self._col = 0
            return
        if tag not in {"td", "th"}:
            return
        if self._row < 0 or self._in_cell:
            raise WorkflowResultError("Table HTML has a malformed cell structure")
        attributes = dict(attrs)
        self._row_span = self._span(attributes, "rowspan")
        self._col_span = self._span(attributes, "colspan")
        if self._row_span * self._col_span > self.max_cells:
            raise InputLimitError("A table span exceeds the configured cell limit")
        while (self._row, self._col) in self._occupied:
            self._col += 1
        last_row = self._row + self._row_span
        last_col = self._col + self._col_span
        if last_row * last_col > self.max_cells:
            raise InputLimitError("A table exceeds the configured cell limit")
        self._text = []
        self._in_cell = True

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() not in {"td", "th"} or not self._in_cell:
            return
        for row in range(self._row, self._row + self._row_span):
            for col in range(self._col, self._col + self._col_span):
                if (row, col) in self._occupied:
                    raise WorkflowResultError("Table HTML contains overlapping spans")
                self._occupied.add((row, col))
                if len(self._occupied) > self.max_cells:
                    raise InputLimitError("A table exceeds the configured cell limit")
        self._cells.append(
            {
                "row": self._row,
                "col": self._col,
                "row_span": self._row_span,
                "col_span": self._col_span,
                "text": " ".join(part.strip() for part in self._text if part.strip()),
            }
        )
        self._col += self._col_span
        self._in_cell = False

    def handle_data(self, data: str) -> None:
        if self._in_cell:
            self._text.append(data)

    @property
    def cells(self) -> list[dict[str, Any]]:
        return self._cells


def _parse_html_cells(html: str, max_cells: int | None = None) -> list[dict[str, Any]]:
    parser = _TableHTMLParser(max_cells or get_settings().max_table_cells)
    parser.feed(html)
    parser.close()
    return parser.cells


def _cells_to_grid(cells: list[dict[str, Any]], max_cells: int | None = None) -> list[list[str]]:
    if not cells:
        return []
    limit = max_cells or get_settings().max_table_cells
    rows = max(cell["row"] + cell["row_span"] for cell in cells)
    columns = max(cell["col"] + cell["col_span"] for cell in cells)
    if rows * columns > limit:
        raise InputLimitError("A table exceeds the configured cell limit")
    grid = [["" for _ in range(columns)] for _ in range(rows)]
    for cell in cells:
        grid[cell["row"]][cell["col"]] = str(cell["text"])
    return grid


def _cell_bbox(raw: Any) -> tuple[float, float, float, float] | None:
    if not isinstance(raw, Sequence) or isinstance(raw, (str, bytes)):
        return None
    values: list[float] = []
    for value in raw:
        if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
            values.extend(float(item) for item in value if isinstance(item, (int, float)))
        elif isinstance(value, (int, float)):
            values.append(float(value))
    if len(values) == 4:
        x1, y1, x2, y2 = values
    elif len(values) >= 8 and len(values) % 2 == 0:
        xs = values[0::2]
        ys = values[1::2]
        x1, y1, x2, y2 = min(xs), min(ys), max(xs), max(ys)
    else:
        return None
    if not all(math.isfinite(value) for value in (x1, y1, x2, y2)) or x1 >= x2 or y1 >= y2:
        return None
    return x1, y1, x2, y2


def _cluster_rows(values: list[float], count: int) -> tuple[list[int], float] | None:
    if count < 1 or count > len(values):
        return None
    ordered = sorted(values)
    if count == 1:
        center = sum(values) / len(values)
        return [0] * len(values), sum((value - center) ** 2 for value in values)
    centers = [ordered[round(index * (len(ordered) - 1) / (count - 1))] for index in range(count)]
    labels = [0] * len(values)
    for _ in range(20):
        labels = [min(range(count), key=lambda index: abs(value - centers[index])) for value in values]
        groups = [[value for value, label in zip(values, labels, strict=True) if label == index] for index in range(count)]
        if any(not group for group in groups):
            return None
        updated = [sum(group) / len(group) for group in groups]
        if max(abs(left - right) for left, right in zip(centers, updated, strict=True)) < 1e-6:
            centers = updated
            break
        centers = updated
    order = {old: new for new, old in enumerate(sorted(range(count), key=lambda index: centers[index]))}
    normalized = [order[label] for label in labels]
    score = sum((value - centers[label]) ** 2 for value, label in zip(values, labels, strict=True))
    return normalized, score


def normalize_table_cells(
    cells: list[dict[str, Any]],
    max_cells: int | None = None,
) -> tuple[list[dict[str, Any]], bool]:
    """Use cell geometry to restore visual row/column order on skewed scans."""
    if not cells:
        return [], False
    limit = max_cells or get_settings().max_table_cells
    copied = [dict(cell) for cell in cells]
    row_count = max(int(cell.get("row", 0)) + int(cell.get("row_span", 1)) for cell in copied)
    column_count = max(int(cell.get("col", 0)) + int(cell.get("col_span", 1)) for cell in copied)
    if row_count * column_count > limit or any(int(cell.get("row_span", 1)) != 1 for cell in copied):
        return copied, False

    boxes = [_cell_bbox(cell.get("bbox")) for cell in copied]
    if any(box is None for box in boxes):
        return copied, False
    centers = [((box[0] + box[2]) / 2, (box[1] + box[3]) / 2) for box in boxes if box]
    best: tuple[list[int], float] | None = None
    for step in range(_SKEW_STEPS):
        slope = -_MAX_TABLE_SKEW + (2 * _MAX_TABLE_SKEW * step / (_SKEW_STEPS - 1))
        adjusted = [y - slope * x for x, y in centers]
        clustered = _cluster_rows(adjusted, row_count)
        if clustered is None:
            continue
        labels, error = clustered
        if best is None or error < best[1]:
            best = labels, error
    if best is None:
        return copied, False

    grouped: list[list[tuple[float, dict[str, Any]]]] = [[] for _ in range(row_count)]
    for cell, (x, _), row in zip(copied, centers, best[0], strict=True):
        grouped[row].append((x, cell))
    normalized: list[dict[str, Any]] = []
    for row, group in enumerate(grouped):
        group.sort(key=lambda item: item[0])
        if sum(int(cell.get("col_span", 1)) for _, cell in group) != column_count:
            return copied, False
        column = 0
        for _, cell in group:
            cell["row"] = row
            cell["col"] = column
            normalized.append(cell)
            column += int(cell.get("col_span", 1))
    _cells_to_grid(normalized, limit)
    return normalized, True


def _geometry_cells_for_mismatched_html(
    cell_boxes: Sequence[Any],
    html_cells: list[dict[str, Any]],
    max_cells: int,
) -> list[dict[str, Any]]:
    """Build a conservative visual grid when HTML and detector cell counts differ.

    Paddle occasionally emits one HTML cell for a span while the cell detector emits
    the physical boxes (or vice versa).  Geometry is authoritative for reading order;
    HTML is used only as text evidence.  The resulting table must still pass primary
    OCR coverage verification before it can be accepted or exported.
    """
    valid = [(raw, _cell_bbox(raw)) for raw in cell_boxes]
    valid = [(raw, box) for raw, box in valid if box is not None]
    if not valid or len(valid) > max_cells:
        return []

    overall = union_bbox([raw for raw, _ in valid])
    if len(overall) == 4:
        width = max(1.0, overall[2] - overall[0])
        height = max(1.0, overall[3] - overall[1])
        aspect = width / height
        box_widths = sorted(box[2] - box[0] for _, box in valid)
        box_heights = sorted(box[3] - box[1] for _, box in valid)
        median_cell_aspect = max(
            0.1,
            box_widths[len(box_widths) // 2] / box_heights[len(box_heights) // 2],
        )
        target_grid_ratio = max(0.1, aspect / median_cell_aspect)
        factor_pairs = [
            (rows, len(valid) // rows)
            for rows in range(1, len(valid) + 1)
            if len(valid) % rows == 0
            and (
                (aspect >= 1.0 and len(valid) // rows >= rows)
                or (aspect < 1.0 and rows >= len(valid) // rows)
            )
        ]
        if factor_pairs:
            row_count, column_count = min(
                factor_pairs,
                key=lambda pair: abs(
                    math.log(max(0.1, (pair[1] / pair[0]) / target_grid_ratio))
                ),
            )
            seeded = [
                {
                    "row": index // column_count,
                    "col": index % column_count,
                    "row_span": 1,
                    "col_span": 1,
                    # Mismatched/transposed HTML is evidence, not safe cell ordering.
                    # High-confidence primary OCR fills these cells after geometry is fixed.
                    "text": "",
                    "bbox": json_safe(raw),
                }
                for index, (raw, _) in enumerate(valid)
            ]
            normalized, succeeded = normalize_table_cells(seeded, max_cells)
            if succeeded and len(normalized) == len(valid):
                return normalized

    heights = sorted(box[3] - box[1] for _, box in valid)
    median_height = heights[len(heights) // 2]
    tolerance = max(4.0, median_height * 0.7)
    ordered = sorted(valid, key=lambda item: ((item[1][1] + item[1][3]) / 2, item[1][0]))
    rows: list[list[tuple[Any, tuple[float, float, float, float]]]] = []
    row_centers: list[float] = []
    for raw, box in ordered:
        center_y = (box[1] + box[3]) / 2
        if not rows or abs(center_y - row_centers[-1]) > tolerance:
            rows.append([(raw, box)])
            row_centers.append(center_y)
        else:
            rows[-1].append((raw, box))
            row_centers[-1] = sum((item[1][1] + item[1][3]) / 2 for item in rows[-1]) / len(
                rows[-1]
            )

    try:
        html_grid = _cells_to_grid(html_cells, max_cells)
    except (InputLimitError, WorkflowResultError):
        html_grid = []
    cells: list[dict[str, Any]] = []
    for row_index, row in enumerate(rows):
        row.sort(key=lambda item: item[1][0])
        for column_index, (raw, _) in enumerate(row):
            text = ""
            if row_index < len(html_grid) and column_index < len(html_grid[row_index]):
                text = html_grid[row_index][column_index]
            cells.append(
                {
                    "row": row_index,
                    "col": column_index,
                    "row_span": 1,
                    "col_span": 1,
                    "text": text,
                    "bbox": json_safe(raw),
                }
            )
    return cells


def table_quality_score(table: Mapping[str, Any]) -> float:
    cells = table.get("cells")
    if not isinstance(cells, Sequence) or isinstance(cells, (str, bytes)) or not cells:
        return 0.0
    valid = [cell for cell in cells if isinstance(cell, Mapping)]
    if not valid:
        return 0.0
    blank_ratio = sum(not str(cell.get("text", "")).strip() for cell in valid) / len(valid)
    box_ratio = sum(_cell_bbox(cell.get("bbox")) is not None for cell in valid) / len(valid)
    normalized_bonus = 0.1 if table.get("geometry_normalized") else 0.0
    has_spans = any(
        int(cell.get("row_span", 1)) > 1 or int(cell.get("col_span", 1)) > 1
        for cell in valid
    )
    geometry_penalty = (
        0.2
        if box_ratio == 1.0 and len(valid) > 1 and not has_spans and not table.get("geometry_normalized")
        else 0.0
    )
    occupied: set[tuple[int, int]] = set()
    overlap = False
    for cell in valid:
        row = int(cell.get("row", 0))
        column = int(cell.get("col", 0))
        row_span = int(cell.get("row_span", 1))
        column_span = int(cell.get("col_span", 1))
        if row < 0 or column < 0 or row_span < 1 or column_span < 1:
            overlap = True
            break
        for current_row in range(row, row + row_span):
            for current_column in range(column, column + column_span):
                position = current_row, current_column
                if position in occupied:
                    overlap = True
                occupied.add(position)
    overlap_penalty = 0.35 if overlap else 0.0
    geometry_score = round(
        max(
            0.0,
            min(
                1.0,
                0.65
                + 0.25 * box_ratio
                + normalized_bonus
                - 0.5 * blank_ratio
                - geometry_penalty
                - overlap_penalty,
            ),
        ),
        4,
    )
    components = table.get("quality_components")
    if not isinstance(components, Mapping):
        return geometry_score
    source_coverage = float(components.get("source_coverage", 0.0) or 0.0)
    text_agreement = float(components.get("text_agreement", 0.0) or 0.0)
    table_token_coverage = float(components.get("table_token_coverage", 0.0) or 0.0)
    orientation = float(components.get("orientation", 0.0) or 0.0)
    bbox_identity = float(components.get("bbox_identity", 0.0) or 0.0)
    plausibility = float(components.get("plausibility", 0.0) or 0.0)
    return round(
        0.20 * geometry_score
        + 0.25 * source_coverage
        + 0.15 * text_agreement
        + 0.10 * table_token_coverage
        + 0.10 * orientation
        + 0.10 * bbox_identity
        + 0.05 * plausibility
        + 0.05 * float(components.get("cell_html_correspondence", 1.0) or 0.0),
        4,
    )


def _token_counter(value: str) -> Counter[str]:
    return Counter(re.findall(r"[\w]+", value.casefold(), flags=re.UNICODE))


def _counter_coverage(source: Counter[str], candidate: Counter[str]) -> float:
    total = sum(source.values())
    if total == 0:
        return 1.0
    matched = sum(min(count, candidate.get(token, 0)) for token, count in source.items())
    return matched / total


def _cell_source_text(
    cell: Mapping[str, Any],
    matches: Sequence[Mapping[str, Any]],
) -> tuple[str, float]:
    ordered = sorted(
        matches,
        key=lambda line: (
            (bbox_values(line.get("normalized_bbox") or line.get("bbox")) or (0, 0, 0, 0))[1],
            (bbox_values(line.get("normalized_bbox") or line.get("bbox")) or (0, 0, 0, 0))[0],
        ),
    )
    text = " ".join(str(line.get("text", "")).strip() for line in ordered).strip()
    confidence = (
        sum(float(line.get("confidence", 0.0) or 0.0) for line in ordered) / len(ordered)
        if ordered
        else 0.0
    )
    return text, confidence


def _assign_lines_to_cells(
    cells: Sequence[Mapping[str, Any]],
    lines: Sequence[Mapping[str, Any]],
) -> dict[int, list[Mapping[str, Any]]]:
    """Assign each primary line to at most one visual cell."""
    assigned: dict[int, list[Mapping[str, Any]]] = {index: [] for index in range(len(cells))}
    for line in lines:
        line_box = bbox_values(line.get("normalized_bbox") or line.get("bbox"))
        if line_box is None:
            continue
        center_x = (line_box[0] + line_box[2]) / 2
        center_y = (line_box[1] + line_box[3]) / 2
        candidates: list[tuple[float, int]] = []
        for index, cell in enumerate(cells):
            cell_box = bbox_values(cell.get("normalized_bbox") or cell.get("bbox"))
            if cell_box is None:
                continue
            overlap = overlap_ratio(line, cell)
            centered = (
                cell_box[0] <= center_x <= cell_box[2]
                and cell_box[1] <= center_y <= cell_box[3]
            )
            if overlap < 0.20 and not centered:
                continue
            cell_width = max(1e-9, cell_box[2] - cell_box[0])
            cell_height = max(1e-9, cell_box[3] - cell_box[1])
            cell_x = (cell_box[0] + cell_box[2]) / 2
            cell_y = (cell_box[1] + cell_box[3]) / 2
            distance = abs(center_x - cell_x) / cell_width + abs(center_y - cell_y) / cell_height
            top_aligned = cell_box[1] <= center_y <= cell_box[1] + 0.35 * cell_height
            score = (
                overlap
                + (0.25 if centered else 0.0)
                + (0.08 if top_aligned else 0.0)
                - 0.05 * distance
            )
            candidates.append((score, index))
        if candidates:
            _, selected = max(candidates)
            assigned[selected].append(line)
    return assigned


def verify_table_against_lines(
    table: dict[str, Any],
    lines: Sequence[Mapping[str, Any]],
    *,
    quality_threshold: float,
    source_coverage_threshold: float = 0.90,
) -> dict[str, Any]:
    """Reconcile a structured grid with primary OCR without inventing cell content."""
    cells = table.get("cells")
    if not isinstance(cells, list) or not cells:
        table["status"] = "needs_review"
        table["quality_score"] = 0.0
        return table

    cell_mappings = [cell for cell in cells if isinstance(cell, Mapping)]
    assigned_lines = _assign_lines_to_cells(cell_mappings, lines)
    agreements: list[float] = []
    for cell_index, cell in enumerate(cells):
        if not isinstance(cell, dict):
            continue
        source_text, source_confidence = _cell_source_text(
            cell,
            assigned_lines.get(cell_index, ()),
        )
        cell["source_text"] = source_text
        cell["source_confidence"] = round(source_confidence, 4)
        structured = str(cell.get("text", "")).strip()
        if source_text and structured:
            agreement = SequenceMatcher(
                None,
                re.sub(r"\W+", "", structured).casefold(),
                re.sub(r"\W+", "", source_text).casefold(),
            ).ratio()
            agreements.append(agreement)
            if source_confidence >= 0.90:
                cell["structure_text"] = structured
                cell["text"] = source_text
                cell["text_source"] = "primary_ocr"
            else:
                cell["text_source"] = "structure"
        elif source_text and not structured and source_confidence >= 0.85:
            cell["text"] = source_text
            cell["text_source"] = "primary_ocr"

    table_box = table.get("normalized_bbox") or table.get("bbox")
    assigned_line_ids = {
        id(line)
        for matches in assigned_lines.values()
        for line in matches
    }
    source_lines = [line for line in lines if id(line) in assigned_line_ids]
    if not any(bbox_values(cell.get("normalized_bbox") or cell.get("bbox")) for cell in cell_mappings):
        source_lines = [line for line in lines if overlap_ratio(line, table_box) >= 0.50]
    source_tokens = _token_counter(" ".join(str(line.get("text", "")) for line in source_lines))
    table_tokens = _token_counter(" ".join(str(cell.get("text", "")) for cell in cells))
    source_coverage = _counter_coverage(source_tokens, table_tokens)
    table_token_coverage = _counter_coverage(table_tokens, source_tokens)
    text_agreement = sum(agreements) / len(agreements) if agreements else source_coverage
    numeric_or_words = [
        token
        for cell in cells
        for token in re.findall(r"\S+", str(cell.get("text", "")))
    ]
    plausible = (
        sum(bool(re.search(r"[\w]", token, flags=re.UNICODE)) for token in numeric_or_words)
        / len(numeric_or_words)
        if numeric_or_words
        else 0.0
    )
    bbox_identity = float(table.get("bbox_identity_score", 0.0) or 0.0)
    if not source_lines:
        bbox_identity = min(bbox_identity, 0.5)
    components = {
        "source_coverage": round(source_coverage, 4),
        "table_token_coverage": round(table_token_coverage, 4),
        "text_agreement": round(text_agreement, 4),
        "orientation": 1.0 if table.get("geometry_normalized") else 0.5,
        "bbox_identity": round(bbox_identity, 4),
        "plausibility": round(plausible, 4),
        "cell_html_correspondence": round(
            float(table.get("cell_html_correspondence", 1.0) or 0.0), 4
        ),
    }
    table["quality_components"] = components
    score = table_quality_score(table)
    table["quality_score"] = score
    accepted = (
        bool(source_lines)
        and score >= quality_threshold
        and source_coverage >= source_coverage_threshold
        and bbox_identity >= 0.25
    )
    table["status"] = "accepted" if accepted else "needs_review"
    table["source_line_ids"] = [
        str(line.get("line_id")) for line in source_lines if line.get("line_id")
    ]
    return table


def _spreadsheet_safe(value: object) -> str:
    text = str(value)
    if text.lstrip().startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


def _save_excel(cells: list[dict[str, Any]], path: Path) -> str:
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table"
    for item in cells:
        row = item["row"] + 1
        column = item["col"] + 1
        cell = worksheet.cell(row=row, column=column)
        cell.value = _spreadsheet_safe(item["text"])
        cell.data_type = "s"
        end_row = row + item["row_span"] - 1
        end_column = column + item["col_span"] - 1
        if end_row > row or end_column > column:
            worksheet.merge_cells(
                start_row=row,
                start_column=column,
                end_row=end_row,
                end_column=end_column,
            )
    temporary = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp.xlsx")
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return str(path)


def _save_csv(grid: list[list[str]], path: Path) -> str:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerows([[_spreadsheet_safe(value) for value in row] for row in grid])
    atomic_write_text(path, stream.getvalue())
    return str(path)


def _table_bbox(blocks: list[Mapping[str, Any]], index: int) -> Any:
    if index < len(blocks):
        return json_safe(blocks[index].get("block_bbox", []))
    return []


def _match_table_blocks(
    candidates: list[dict[str, Any]],
    blocks: list[Mapping[str, Any]],
) -> None:
    """Associate tables spatially; list order is not a stable PaddleOCR contract."""
    available = set(range(len(blocks)))
    for candidate in sorted(
        candidates,
        key=lambda item: bool(bbox_values(item.get("derived_bbox"))),
        reverse=True,
    ):
        derived = candidate.get("derived_bbox")
        scored = sorted(
            (
                (bbox_iou(derived, blocks[index].get("block_bbox", [])), index)
                for index in available
            ),
            reverse=True,
        )
        if scored and scored[0][0] > 0.0:
            score, selected = scored[0]
            available.remove(selected)
            candidate["bbox"] = json_safe(blocks[selected].get("block_bbox", []))
            candidate["bbox_identity_score"] = round(score, 4)
            candidate["bbox_match"] = "spatial"
        elif len(candidates) == 1 and len(blocks) == 1:
            candidate["bbox"] = json_safe(blocks[0].get("block_bbox", []))
            candidate["bbox_identity_score"] = 0.5
            candidate["bbox_match"] = "single_unambiguous"
            available.discard(0)
        else:
            candidate["bbox"] = json_safe(derived)
            candidate["bbox_identity_score"] = 0.0
            candidate["bbox_match"] = "unmatched"


def extract_tables_from_prediction(
    prediction: object,
    *,
    page_index: int,
    first_table_index: int,
) -> list[dict[str, Any]]:
    settings = get_settings()
    payload = result_payload(prediction, workflow="Table extraction")
    serialized_tables = payload.get("table_res_list")
    parsing_blocks = payload.get("parsing_res_list")
    if not isinstance(parsing_blocks, Sequence) or isinstance(parsing_blocks, (str, bytes)):
        raise WorkflowResultError(
            "Table extraction returned an unsupported PaddleOCR result schema: parsing_res_list is missing"
        )
    if any(not isinstance(block, Mapping) for block in parsing_blocks):
        raise WorkflowResultError("Table extraction returned a malformed parsing block")
    table_blocks = [
        block
        for block in parsing_blocks
        if isinstance(block, Mapping) and str(block.get("block_label", "")).lower() == "table"
    ]
    if not isinstance(serialized_tables, Sequence) or isinstance(serialized_tables, (str, bytes)):
        # PP-Structure omits table_res_list on pages where no table was detected.
        # That is a valid empty result, not an upstream schema change.
        if not table_blocks:
            return []
        raise WorkflowResultError(
            "Table extraction returned an unsupported PaddleOCR result schema: table_res_list is missing"
        )
    direct_tables = prediction.get("table_res_list") if hasattr(prediction, "get") else None
    if not serialized_tables and direct_tables is None:
        return []
    if not isinstance(direct_tables, Sequence) or isinstance(direct_tables, (str, bytes)):
        raise WorkflowResultError("Table extraction returned an unsupported live table_res_list")
    if len(direct_tables) != len(serialized_tables):
        raise WorkflowResultError("Table extraction returned inconsistent table results")

    candidates: list[dict[str, Any]] = []
    for local_index, table_result in enumerate(direct_tables):
        table_payload = result_payload(table_result, workflow="Table extraction")
        html = table_payload.get("pred_html")
        cell_boxes = table_payload.get("cell_box_list", [])
        if not isinstance(html, str):
            raise WorkflowResultError("Table extraction returned a table without pred_html")
        html_cells = _parse_html_cells(html, settings.max_table_cells)
        if not isinstance(cell_boxes, Sequence) or isinstance(cell_boxes, (str, bytes)):
            raise WorkflowResultError("Table extraction returned malformed cell_box_list")
        correspondence = (
            min(len(cell_boxes), len(html_cells)) / max(len(cell_boxes), len(html_cells))
            if cell_boxes or html_cells
            else 1.0
        )
        mismatch_warning = None
        geometry_from_mismatch = False
        if len(cell_boxes) == len(html_cells):
            cells = html_cells
            for cell, bbox in zip(cells, cell_boxes, strict=True):
                cell["bbox"] = json_safe(bbox)
        elif cell_boxes:
            cells = _geometry_cells_for_mismatched_html(
                cell_boxes,
                html_cells,
                settings.max_table_cells,
            )
            if not cells:
                cells = html_cells
                for cell in cells:
                    cell["bbox"] = []
            else:
                geometry_from_mismatch = True
            mismatch_warning = (
                f"Structure returned {len(cell_boxes)} cell boxes for "
                f"{len(html_cells)} HTML cells; visual geometry was used conservatively."
            )
        else:
            cells = html_cells
            for cell in cells:
                cell["bbox"] = []
        normalized_cells, geometry_normalized = normalize_table_cells(
            cells,
            settings.max_table_cells,
        )
        geometry_normalized = geometry_normalized or geometry_from_mismatch
        derived_bbox = union_bbox([cell.get("bbox", []) for cell in normalized_cells])
        table = {
            "table_index": first_table_index + local_index,
            "page_index": page_index,
            "bbox": derived_bbox,
            "derived_bbox": derived_bbox,
            "html": html,
            "cells": normalized_cells,
            "geometry_normalized": geometry_normalized,
            "cell_html_correspondence": round(correspondence, 4),
            "structure_warning": mismatch_warning,
            "orientation": 0,
            "retry_metadata": {"attempted": False, "accepted": False},
            "excel_path": None,
            "csv_path": None,
        }
        candidates.append(table)
    _match_table_blocks(candidates, table_blocks)
    for table in candidates:
        table["quality_score"] = table_quality_score(table)
        table["status"] = "unverified"
    return candidates


def run_table_extraction(
    image_input: str | Path,
    save_output: bool = True,
    output_dir: str | Path | None = None,
    *,
    job_id: str | None = None,
    progress_callback: ProgressCallback | None = None,
    language: str | None = None,
    document_profile: str = "auto",
) -> dict[str, Any]:
    settings = get_settings()
    selected_language = validate_language(
        language,
        ocr_version=settings.ocr_version,
        structure_version=settings.structure_ocr_version,
        default=settings.ocr_lang,
    )
    selected_profile = validate_document_profile(document_profile)
    input_path = str(image_input)
    tables: list[dict[str, Any]] = []
    elements: list[dict[str, Any]] = []
    pages: list[dict[str, Any]] = []
    page_count = 0

    def consume(
        predictions: object,
        page_index: int,
        *,
        width: int | None = None,
        height: int | None = None,
        render_scale: float | None = None,
    ) -> None:
        nonlocal page_count
        result_count = 0
        for result in predictions:
            result_count += 1
            payload = result_payload(result, workflow="Table extraction")
            parsed = parsing_elements(
                payload,
                workflow="Table extraction",
                page_index=page_index,
                first_index=len(elements),
            )
            page_tables = extract_tables_from_prediction(
                result,
                page_index=page_index,
                first_table_index=len(tables),
            )
            if width and height:
                for element in parsed:
                    element["normalized_bbox"] = normalized_bbox(
                        element.get("bbox"), width, height
                    )
                for table in page_tables:
                    table["normalized_bbox"] = normalized_bbox(
                        table.get("bbox"), width, height
                    )
                    for cell in table.get("cells", []):
                        if isinstance(cell, dict):
                            cell["normalized_bbox"] = normalized_bbox(
                                cell.get("bbox"), width, height
                            )
            elements.extend(parsed)
            tables.extend(page_tables)
            page_count = max(page_count, page_index + 1)
        if result_count != 1:
            raise WorkflowResultError(
                "Table extraction returned an unexpected number of results for one page"
            )
        pages.append(
            {
                "page_index": page_index,
                "page_number": page_index + 1,
                "width": width,
                "height": height,
                "render_scale": render_scale,
            }
        )

    emit_progress(
        progress_callback,
        stage="initializing",
        message="Loading the table extraction engine",
        completed_pages=0,
    )
    try:
        with get_manager().session(
            "structure", language=selected_language, document_profile=selected_profile
        ) as engine:
            if is_pdf(input_path):
                for page in iter_pdf_pages(
                    input_path,
                    settings=settings,
                    progress_callback=progress_callback,
                ):
                    logger.info("Table PDF page %d/%d", page.index + 1, page.count)
                    emit_progress(
                        progress_callback,
                        stage="table_extraction",
                        message=f"Extracting tables from page {page.number} of {page.count}",
                        current_page=page.number,
                        total_pages=page.count,
                        completed_pages=page.index,
                    )
                    consume(
                        engine.predict(page.image),
                        page.index,
                        width=page.image.shape[1],
                        height=page.image.shape[0],
                        render_scale=page.render_scale,
                    )
                    emit_progress(
                        progress_callback,
                        stage="page_complete",
                        message=f"Finished page {page.number} of {page.count}",
                        current_page=page.number,
                        total_pages=page.count,
                        completed_pages=page.number,
                    )
            else:
                emit_progress(
                    progress_callback,
                    stage="table_extraction",
                    message="Extracting tables from image",
                    current_page=1,
                    total_pages=1,
                    completed_pages=0,
                )
                with Image.open(input_path) as image:
                    width, height = image.size
                consume(engine.predict(input_path), 0, width=width, height=height)
                emit_progress(
                    progress_callback,
                    stage="page_complete",
                    message="Finished image",
                    current_page=1,
                    total_pages=1,
                    completed_pages=1,
                )
    except Exception as exc:
        translated = translate_inference_error(exc)
        if translated is not exc:
            raise translated from exc
        raise
    finally:
        if settings.clear_gpu_cache_after_document:
            clear_gpu_cache()

    warnings: list[str] = []
    primary_result: dict[str, Any] = {}
    try:
        from app.workflows.ocr_text import run_text_recognition

        emit_progress(
            progress_callback,
            stage="quality_validation",
            message="Verifying table cells against primary text recognition",
        )
        primary_result = run_text_recognition(
            input_path,
            progress_callback=progress_callback,
            language=selected_language,
            document_profile=selected_profile,
        )
        primary_lines = primary_result.get("lines", [])
        for table in tables:
            page_lines = [
                line
                for line in primary_lines
                if isinstance(line, Mapping)
                and int(line.get("page_index", 0) or 0)
                == int(table.get("page_index", 0) or 0)
            ]
            verify_table_against_lines(
                table,
                page_lines,
                quality_threshold=settings.table_quality_threshold,
            )
            if table.get("status") != "accepted":
                warnings.append(
                    f"Table {int(table['table_index']) + 1} could not be fully verified "
                    "against primary OCR and needs review."
                )
    except Exception:
        logger.warning("Primary OCR table verification failed", exc_info=True)
        warnings.append(
            "Primary OCR table verification was unavailable; table exports need review."
        )

    emit_progress(
        progress_callback,
        stage="finalizing",
        message="Writing table exports",
        completed_pages=1 if not is_pdf(input_path) else None,
    )
    if save_output:
        directory = resolve_output_directory(output_dir, job_id, settings)
        directory.mkdir(parents=True, exist_ok=True)
        for table in tables:
            if table.get("status") != "accepted":
                continue
            index = table["table_index"]
            cells = table["cells"]
            grid = _cells_to_grid(cells, settings.max_table_cells)
            if grid:
                excel_path = Path(_save_excel(cells, directory / f"table_{index}.xlsx"))
                csv_path = Path(_save_csv(grid, directory / f"table_{index}.csv"))
                table["excel_path"] = public_path(excel_path, settings.project_root)
                table["csv_path"] = public_path(csv_path, settings.project_root)
        atomic_write_json(directory / "tables.json", tables)
    return {
        "tables": tables,
        "table_count": len(tables),
        "elements": elements,
        "page_count": page_count or 1,
        "language": selected_language,
        "requested_document_profile": selected_profile,
        "lines": primary_result.get("lines", []),
        "full_text": primary_result.get("full_text", ""),
        "adaptive_retry_attempted": primary_result.get("adaptive_retry_attempted", 0),
        "adaptive_retry_accepted": primary_result.get("adaptive_retry_accepted", 0),
        "warnings": warnings,
        "pages": pages,
    }
